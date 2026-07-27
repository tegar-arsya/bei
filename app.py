import io
import json
import os
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlencode

import numpy as np
import pandas as pd
import requests
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analytics.ai import build_runtime_context, call_openrouter_structured
from analytics.chart import create_ai_chart
from analytics.goapi import (
    GoAPIError,
    clear_goapi_cache,
    fetch_goapi_broker_summary,
    fetch_goapi_historical,
    fetch_goapi_snapshot,
    goapi_is_configured,
)
from analytics.quant import build_price_history, build_quant_context
from analytics.validator import validate_analysis
from storage.persistence import (
    IMPORT_MIGRATION_FILE,
    load_import_bundle,
    load_import_dates,
    save_analysis,
    save_bei_import,
    supabase_is_configured,
)

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass


st.set_page_config(page_title="BEI Stock Screener", layout="wide")


REQUIRED_COLUMNS = {
    "ringkasan_saham": [
        "kode saham", "nama perusahaan", "sebelumnya", "open price",
        "tanggal perdagangan terakhir", "first trade", "tertinggi", "terendah",
        "penutupan", "selisih", "volume", "nilai", "frekuensi",
        "offer", "offer volume", "bid", "bid volume", "listed shares",
        "tradeble shares", "weight for index", "foreign sell", "foreign buy",
        "non regular volume", "non regular value", "non regular frequency",
    ],
    "ringkasan_broker": ["kode perusahaan", "nama perusahaan", "volume", "nilai", "frekuensi"],
    "ringkasan_perdagangan": ["id instrument", "id board", "volume", "nilai", "frekuensi"],
    "daftar_saham": ["kode", "nama perusahaan", "tanggal pencatatan", "saham", "papan pencatatan"],
}
ALTERNATE_COLUMNS = {
    "daftar_saham": ["id instrument", "id board", "volume", "nilai", "frekuensi"],
}
INDEX_REQUIRED_COLUMNS = [
    "kode indeks", "sebelumnya", "tertinggi", "terendah",
    "penutupan", "selisih", "volume", "nilai", "frekuensi",
]
DEFAULT_WEIGHTS = {
    "momentum": 0.25,
    "liquidity": 0.20,
    "flow": 0.20,
    "market_activity": 0.15,
    "volume_trend": 0.10,
    "price_structure": 0.05,
    "broker": 0.05,
}
DEFAULT_ALERT_RULES = {
    "stock_min_score": 72.0,
    "stock_min_change": 3.0,
    "stock_dump_change": -4.0,
    "stock_min_rel_volume": 1.5,
}

TV_FIELDS = [
    "name", "close", "change", "volume",
    "RSI", "RSI[1]",
    "MACD.macd", "MACD.signal",
    "BB.upper", "BB.lower", "BB.basis",
    "EMA20", "EMA50", "EMA200",
    "Stoch.K", "Stoch.D",
    "ADX", "ADX+DI", "ADX-DI",
    "Recommend.All", "Recommend.MA", "Recommend.Other",
    "CCI20", "Perf.W", "Perf.1M",
    "relative_volume_10d_calc", "Mom", "AO",
]
STOCK_AUTO_FIELDS = [
    "name", "description", "close", "change", "volume", "Value.Traded", "market_cap_basic",
    "RSI", "RSI[1]",
    "MACD.macd", "MACD.signal",
    "BB.upper", "BB.lower", "BB.basis",
    "EMA20", "EMA50", "EMA200",
    "Stoch.K", "Stoch.D",
    "ADX", "ADX+DI", "ADX-DI",
    "Recommend.All", "Recommend.MA", "Recommend.Other",
    "CCI20", "Perf.W", "Perf.1M",
    "relative_volume_10d_calc", "Mom", "AO",
    "sector", "type",
]
STOCK_SORT_OPTIONS = {
    "Nilai transaksi": "Value.Traded",
    "Volume": "volume",
    "Momentum harian": "change",
    "Performa 1 minggu": "Perf.W",
    "Performa 1 bulan": "Perf.1M",
    "Market cap": "market_cap_basic",
}
APP_PAGES = ["Home", "Saham BEI", "Watchlist & Alerts"]

DATA_DIR = "data"
WATCHLIST_FILE = os.path.join(DATA_DIR, "stock_watchlist.json")
HISTORY_FILE = os.path.join(DATA_DIR, "stock_history.json")
ALERT_RULES_FILE = os.path.join(DATA_DIR, "stock_alert_rules.json")
MAX_HISTORY_ROWS = 5000


def tv_rec_label(value):
    try:
        if value is None:
            return "N/A"
        number = float(value)
        if np.isnan(number):
            return "N/A"
        if number >= 0.5:
            return "Strong Buy"
        if number >= 0.1:
            return "Buy"
        if number >= -0.1:
            return "Neutral"
        if number >= -0.5:
            return "Sell"
        return "Strong Sell"
    except Exception:
        return "N/A"


def tv_rec_emoji(label):
    return {
        "Strong Buy": "🟢🟢",
        "Buy": "🟢",
        "Neutral": "🟡",
        "Sell": "🔴",
        "Strong Sell": "🔴🔴",
    }.get(label, "⚪")


def normalize_column_name(name):
    return " ".join(str(name).strip().lower().replace("_", " ").split())


def normalize_columns(df):
    return df.rename(columns={column: normalize_column_name(column) for column in df.columns})


def find_col(df, target):
    normalized = normalize_column_name(target)
    for column in df.columns:
        if normalize_column_name(column) == normalized:
            return column
    raise KeyError(f"Kolom '{target}' tidak ditemukan")


def to_numeric(series):
    if pd.api.types.is_numeric_dtype(series):
        return series.astype(float)
    raw = series.astype(str).str.strip().replace({"": np.nan, "-": np.nan, "--": np.nan})
    raw = raw.str.replace("%", "", regex=False).str.replace(" ", "", regex=False)
    parsed_plain = pd.to_numeric(raw, errors="coerce")
    parsed_id = pd.to_numeric(
        raw.str.replace(".", "", regex=False).str.replace(",", ".", regex=False),
        errors="coerce",
    )
    parsed_comma = pd.to_numeric(raw.str.replace(",", "", regex=False), errors="coerce")
    return parsed_plain.fillna(parsed_id).fillna(parsed_comma)


def percentile_series(series):
    clean = series.copy().replace([np.inf, -np.inf], np.nan)
    if clean.notna().sum() <= 1:
        return pd.Series([50.0] * len(clean), index=clean.index)
    return clean.rank(pct=True) * 100


def safe_col(df, column):
    return to_numeric(df[find_col(df, column)])


def safe_num(value, default=0.0):
    try:
        return default if pd.isna(value) else float(value)
    except Exception:
        return default


def has_columns(df, columns):
    existing = {normalize_column_name(column) for column in df.columns}
    return {normalize_column_name(column) for column in columns}.issubset(existing)


def validate_columns(df, required_columns):
    existing = [normalize_column_name(column) for column in df.columns]
    missing = [column for column in required_columns if column not in existing]
    return len(missing) == 0, missing


def color_signal(value):
    if isinstance(value, (float, int, np.floating)):
        if value >= 75:
            return "background-color:#1a472a;color:#90ee90"
        if value >= 60:
            return "background-color:#1e3a5f;color:#87ceeb"
        if value >= 45:
            return "background-color:#3d2b00;color:#ffd700"
        return "background-color:#3d0000;color:#ff9999"
    return ""


def render_df_with_style_fallback(df, subset_cols):
    try:
        st.dataframe(df.style.map(color_signal, subset=subset_cols), width="stretch")
    except Exception:
        st.dataframe(df, width="stretch")


def factor_label(score):
    if pd.isna(score):
        return "N/A"
    value = float(score)
    if value >= 75:
        return "Sangat Kuat"
    if value >= 60:
        return "Kuat"
    if value >= 45:
        return "Netral"
    return "Lemah"


def format_compact(value, prefix=""):
    try:
        number = float(value)
        if pd.isna(number):
            return "N/A"
    except Exception:
        return "N/A"
    sign = "-" if number < 0 else ""
    absolute = abs(number)
    for size, suffix in [(1_000_000_000, "B"), (1_000_000, "M"), (1_000, "K")]:
        if absolute >= size:
            return f"{sign}{prefix}{absolute / size:.2f}{suffix}"
    return f"{sign}{prefix}{absolute:.2f}"


def price_text(value):
    try:
        number = float(value)
        if pd.isna(number) or np.isinf(number):
            return "N/A"
        if abs(number) >= 100:
            return f"{number:,.2f}"
        if abs(number) >= 1:
            return f"{number:,.4f}"
        return f"{number:,.8g}"
    except Exception:
        return "N/A"


def format_idr(value, compact=True):
    if compact:
        return format_compact(value, "Rp")
    raw = price_text(value)
    return raw if raw == "N/A" else f"Rp{raw}"


def pct_text(value):
    try:
        number = float(value)
        return "N/A" if pd.isna(number) else f"{number:+.2f}%"
    except Exception:
        return "N/A"


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel_sheet_name(name, used_names):
    """Return a valid, unique Excel worksheet name."""
    base = " ".join(str(name or "Data").split()) or "Data"
    base = base.replace("/", "-").replace("\\", "-")
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate in used_names:
        suffix_text = f" ({suffix})"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def dataframe_to_excel_bytes(frames, chart_sheet=None):
    """Create a readable Excel workbook for Streamlit download buttons.

    The first row is frozen and filtered on every sheet so a large result can
    be inspected directly in Excel. When a quant-history sheet is supplied,
    a native close/EMA chart is also added beside its rows.
    """
    if isinstance(frames, pd.DataFrame):
        frames = {"Data": frames}
    frames = frames or {}
    output = io.BytesIO()
    used_names = set()
    with pd.ExcelWriter(output, engine="openpyxl", date_format="yyyy-mm-dd", datetime_format="yyyy-mm-dd") as writer:
        sheet_map = {}
        for requested_name, frame in frames.items():
            sheet_name = _excel_sheet_name(requested_name, used_names)
            sheet_map[requested_name] = sheet_name
            data = frame.copy() if isinstance(frame, pd.DataFrame) else pd.DataFrame(frame)
            for column in data.columns:
                if isinstance(data[column].dtype, pd.CategoricalDtype):
                    data[column] = data[column].astype(object)
            data.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            worksheet.sheet_view.showGridLines = False
            worksheet.freeze_panes = "A2"
            if worksheet.max_column and worksheet.max_row > 1:
                worksheet.auto_filter.ref = worksheet.dimensions

            header_fill = PatternFill("solid", fgColor="17365D")
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            worksheet.row_dimensions[1].height = 30

            for column_cells in worksheet.columns:
                column_index = column_cells[0].column
                values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:250]]
                longest = max([len(value) for value in values] + [8])
                worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 11), 42)
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top")

        if chart_sheet in sheet_map:
            worksheet = writer.sheets[sheet_map[chart_sheet]]
            headers = {str(cell.value): cell.column for cell in worksheet[1] if cell.value}
            chart_columns = [name for name in ("close", "ema20", "ema50", "ema200") if name in headers]
            if "date" in headers and chart_columns and worksheet.max_row > 1:
                chart = LineChart()
                chart.title = "AI Quant Chart — Close dan EMA"
                chart.y_axis.title = "Harga"
                chart.x_axis.title = "Tanggal / observasi"
                chart.height = 8
                chart.width = 16
                for column_name in chart_columns:
                    values = Reference(
                        worksheet,
                        min_col=headers[column_name],
                        min_row=1,
                        max_row=worksheet.max_row,
                    )
                    chart.add_data(values, titles_from_data=True)
                categories = Reference(
                    worksheet,
                    min_col=headers["date"],
                    min_row=2,
                    max_row=worksheet.max_row,
                )
                chart.set_categories(categories)
                chart.style = 13
                worksheet.add_chart(chart, "L2")
    output.seek(0)
    return output.getvalue()


def build_analysis_excel_bytes(analysis, quant_context=None, chart_history=None):
    """Build a detailed, row-readable workbook for one AI analysis."""
    analysis = analysis if isinstance(analysis, dict) else {}
    quant_context = quant_context if isinstance(quant_context, dict) else {}
    quality = analysis.get("data_quality", {}) or {}
    structure = analysis.get("market_structure", {}) or {}
    levels = analysis.get("levels", {}) or {}
    indicators = quant_context.get("indicators", {}) or {}
    execution = quant_context.get("execution", {}) or {}

    overview = pd.DataFrame([
        {"Bagian": "Ticker", "Nilai": analysis.get("ticker", "")},
        {"Bagian": "Tanggal analisis", "Nilai": analysis.get("as_of_date", "")},
        {"Bagian": "Verdict", "Nilai": analysis.get("verdict", "")},
        {"Bagian": "Confidence", "Nilai": analysis.get("confidence", "")},
        {"Bagian": "Data quality", "Nilai": quality.get("status", "")},
        {"Bagian": "Ringkasan", "Nilai": analysis.get("summary", "")},
        {"Bagian": "Kesimpulan", "Nilai": analysis.get("conclusion", "")},
        {"Bagian": "Disclaimer", "Nilai": analysis.get("disclaimer", "")},
    ])
    structure_rows = [
        {"Komponen": "Trend", "Detail": structure.get("trend", "")},
        {"Komponen": "Momentum", "Detail": structure.get("momentum", "")},
        {"Komponen": "Flow", "Detail": structure.get("flow", "")},
        {"Komponen": "Technical", "Detail": structure.get("technical", "")},
    ]
    quant_rows = [{"Indikator": key, "Nilai": value} for key, value in indicators.items()]
    quant_rows += [{"Indikator": f"execution_{key}", "Nilai": value} for key, value in execution.items()]
    quant_rows += [{"Indikator": "observations", "Nilai": quant_context.get("observations", 0)},
                   {"Indikator": "status", "Nilai": quant_context.get("status", "") }]

    level_rows = []
    for key in ("support", "resistance"):
        for index, value in enumerate(levels.get(key, []) or [], start=1):
            level_rows.append({"Kelompok": key, "Nama": f"{key}_{index}", "Harga": value})
    for key in ("entry_low", "entry_high", "stop", "tp1", "tp2"):
        if key in levels:
            level_rows.append({"Kelompok": "execution", "Nama": key, "Harga": levels.get(key)})
    for item in levels.get("fibonacci", []) or []:
        if isinstance(item, dict):
            level_rows.append({"Kelompok": "fibonacci", "Nama": item.get("label", "Fib"), "Harga": item.get("price"), "Rasio": item.get("ratio")})

    scenario_rows = [
        {
            "Skenario": item.get("name", ""),
            "Kondisi": item.get("condition", ""),
            "Trigger": item.get("trigger", ""),
            "Invalidasi": item.get("invalidation", ""),
            "Probabilitas": item.get("probability", ""),
            "Aksi": item.get("action", ""),
        }
        for item in analysis.get("scenarios", []) or [] if isinstance(item, dict)
    ]
    evidence_rows = []
    for category, values in [
        ("Evidence", structure.get("evidence", [])),
        ("Risk flag", analysis.get("risk_flags", [])),
        ("Monitoring trigger", analysis.get("monitoring_triggers", [])),
        ("Data missing", quality.get("missing", [])),
        ("Data conflict", quality.get("conflicts", [])),
    ]:
        evidence_rows.extend({"Kategori": category, "Detail": value} for value in (values or []))

    drawing_rows = [item for item in analysis.get("drawings", []) or [] if isinstance(item, dict)]
    drawing_frame = pd.DataFrame(drawing_rows, columns=["type", "name", "purpose", "x_start", "x_end", "y_start", "y_end", "evidence"])
    quant_history = pd.DataFrame(quant_context.get("history", []) or [])
    if quant_history.empty and isinstance(chart_history, pd.DataFrame):
        quant_history = chart_history.copy()

    frames = {
        "Ringkasan": overview,
        "Market Structure": pd.DataFrame(structure_rows),
        "Quant Indicators": pd.DataFrame(quant_rows),
        "Levels": pd.DataFrame(level_rows),
        "Scenarios": pd.DataFrame(scenario_rows),
        "Evidence": pd.DataFrame(evidence_rows),
        "Drawings": drawing_frame,
        "AI Quant Chart": quant_history,
    }
    return dataframe_to_excel_bytes(frames, chart_sheet="AI Quant Chart")


def ratio_text(value, suffix="x"):
    try:
        number = float(value)
        if pd.isna(number) or np.isinf(number):
            return "N/A"
        return f"{number:.2f}{suffix}"
    except Exception:
        return "N/A"


def utc_now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def get_static_openrouter_key():
    key = ""
    try:
        key = str(st.secrets.get("OPENROUTER_API_KEY", "")).strip()
    except Exception:
        pass
    return key or os.getenv("OPENROUTER_API_KEY", "").strip()


def get_openrouter_model():
    model = ""
    try:
        model = str(st.secrets.get("OPENROUTER_MODEL", "")).strip()
    except Exception:
        pass
    return model or os.getenv("OPENROUTER_MODEL", "").strip() or "openrouter/auto"


# ── TRADINGVIEW ──────────────────────────────────────────────────────────────
@st.cache_data(ttl=300, show_spinner=False)
def fetch_tv_data(ticker):
    symbol = f"IDX:{ticker.upper().strip()}"
    payload = {
        "symbols": {"tickers": [symbol], "query": {"types": []}},
        "columns": TV_FIELDS,
    }
    headers = {
        "Content-Type": "application/json",
        "Origin": "https://www.tradingview.com",
        "Referer": "https://www.tradingview.com/",
        "User-Agent": "Mozilla/5.0",
    }
    try:
        response = requests.post(
            "https://scanner.tradingview.com/indonesia/scan",
            json=payload,
            headers=headers,
            timeout=15,
        )
        response.raise_for_status()
        data = response.json().get("data", [])
        if not data:
            return None
        values = data[0].get("d", [])
        result = {field: values[index] for index, field in enumerate(TV_FIELDS) if index < len(values)}
        result["symbol"] = symbol
        return result
    except Exception:
        return None


def parse_tv_data(tv):
    if not tv:
        return {}
    close = safe_num(tv.get("close"), np.nan)
    ema20 = safe_num(tv.get("EMA20"), np.nan)
    ema50 = safe_num(tv.get("EMA50"), np.nan)
    ema200 = safe_num(tv.get("EMA200"), np.nan)
    bb_upper = safe_num(tv.get("BB.upper"), np.nan)
    bb_lower = safe_num(tv.get("BB.lower"), np.nan)
    bb_basis = safe_num(tv.get("BB.basis"), np.nan)
    macd = safe_num(tv.get("MACD.macd"), np.nan)
    macd_signal = safe_num(tv.get("MACD.signal"), np.nan)
    bb_position = np.nan
    if not any(pd.isna(value) for value in [close, bb_lower, bb_upper]) and bb_upper > bb_lower:
        bb_position = (close - bb_lower) / (bb_upper - bb_lower) * 100

    def above(value):
        return close > value if not any(pd.isna([close, value])) else None

    above_ema20 = above(ema20)
    above_ema50 = above(ema50)
    above_ema200 = above(ema200)
    ema_count = sum(1 for item in [above_ema20, above_ema50, above_ema200] if item is True)
    macd_cross = None
    if not any(pd.isna([macd, macd_signal])):
        macd_cross = "bullish" if macd > macd_signal else "bearish"
    return {
        "close": close,
        "change_pct": safe_num(tv.get("change"), np.nan),
        "rec_all": tv_rec_label(tv.get("Recommend.All")),
        "rec_ma": tv_rec_label(tv.get("Recommend.MA")),
        "rec_other": tv_rec_label(tv.get("Recommend.Other")),
        "rsi": safe_num(tv.get("RSI"), np.nan),
        "rsi_prev": safe_num(tv.get("RSI[1]"), np.nan),
        "macd": macd,
        "macd_signal": macd_signal,
        "macd_cross": macd_cross,
        "stoch_k": safe_num(tv.get("Stoch.K"), np.nan),
        "stoch_d": safe_num(tv.get("Stoch.D"), np.nan),
        "adx": safe_num(tv.get("ADX"), np.nan),
        "adx_plus": safe_num(tv.get("ADX+DI"), np.nan),
        "adx_minus": safe_num(tv.get("ADX-DI"), np.nan),
        "cci": safe_num(tv.get("CCI20"), np.nan),
        "mom": safe_num(tv.get("Mom"), np.nan),
        "ao": safe_num(tv.get("AO"), np.nan),
        "ema20": ema20,
        "ema50": ema50,
        "ema200": ema200,
        "bb_upper": bb_upper,
        "bb_lower": bb_lower,
        "bb_basis": bb_basis,
        "bb_position": bb_position,
        "ema_trend_count": ema_count,
        "above_ema20": above_ema20,
        "above_ema50": above_ema50,
        "above_ema200": above_ema200,
        "rel_volume": safe_num(tv.get("relative_volume_10d_calc"), np.nan),
        "perf_week": safe_num(tv.get("Perf.W"), np.nan),
        "perf_month": safe_num(tv.get("Perf.1M"), np.nan),
    }


def build_prompt(ticker, company, bei_row, tv, market_regime, multiday_context):
    close = safe_num(bei_row.get("penutupan_num"), np.nan)
    change = safe_num(bei_row.get("change_pct"), np.nan)
    foreign_net = safe_num(bei_row.get("foreign_net"), 0)
    foreign_ratio = safe_num(bei_row.get("foreign_net_ratio"), 0)
    bid_pressure = safe_num(bei_row.get("bid_offer_pressure"), 0)
    atr_pct = max(safe_num(bei_row.get("atr_pct"), 2.0), 0.5)
    signal = safe_num(bei_row.get("signal_strength", bei_row.get("final_score", 0)), 0)
    final_score = safe_num(bei_row.get("final_score"), 0)
    atr_abs = close * atr_pct / 100 if not pd.isna(close) else 0

    def fmt(value, decimals=2):
        return f"{value:.{decimals}f}" if not pd.isna(value) else "N/A"

    multiday_text = ""
    if multiday_context and multiday_context.get("days_data", 1) > 1:
        multiday_text = f"""
MULTI-HARI ({multiday_context["days_data"]} hari):
- Tren skor: {multiday_context["trend_slope"]:+.1f} poin/hari
- Konsistensi score >=60: {multiday_context["score_consistency"]:.0f}%
- Signal: {bei_row.get("signal_label", "-")}"""

    regime_text = ""
    if market_regime:
        regime_text = (
            f"MARKET REGIME: {market_regime['regime_label']} "
            f"(score={market_regime['regime_score']:.1f}, "
            f"IHSG {market_regime['index_change_pct']:+.2f}%)"
        )

    if tv:
        ema_desc = {
            3: "BULLISH - di atas semua EMA 20/50/200",
            2: "CUKUP BULLISH - di atas 2 dari 3 EMA",
            1: "LEMAH - di atas 1 dari 3 EMA",
            0: "BEARISH - di bawah semua EMA",
        }
        rsi_note = "OVERBOUGHT" if tv.get("rsi", 50) > 70 else "OVERSOLD" if tv.get("rsi", 50) < 30 else "normal"
        adx_note = "tren kuat" if tv.get("adx", 0) > 25 else "sideways/lemah"
        tv_text = f"""
TEKNIKAL TRADINGVIEW:
- Rekomendasi: {tv.get("rec_all", "N/A")} (MA: {tv.get("rec_ma", "N/A")} | Osc: {tv.get("rec_other", "N/A")})
- RSI(14): {fmt(tv.get("rsi"), 1)} [{rsi_note}]
- MACD: {tv.get("macd_cross", "N/A")} | MACD={fmt(tv.get("macd"), 4)} | Signal={fmt(tv.get("macd_signal"), 4)}
- Stochastic K/D: {fmt(tv.get("stoch_k"), 1)}/{fmt(tv.get("stoch_d"), 1)}
- ADX: {fmt(tv.get("adx"), 1)} [{adx_note}] | +DI={fmt(tv.get("adx_plus"), 1)} vs -DI={fmt(tv.get("adx_minus"), 1)}
- CCI={fmt(tv.get("cci"), 1)} | Momentum={fmt(tv.get("mom"), 4)} | AO={fmt(tv.get("ao"), 4)}
- EMA: {ema_desc.get(tv.get("ema_trend_count"), "N/A")}
  EMA20={fmt(tv.get("ema20"))} | EMA50={fmt(tv.get("ema50"))} | EMA200={fmt(tv.get("ema200"))}
- Bollinger: {fmt(tv.get("bb_position"), 1)}% | Lower={fmt(tv.get("bb_lower"))} | Mid={fmt(tv.get("bb_basis"))} | Upper={fmt(tv.get("bb_upper"))}
- Relative volume: {fmt(tv.get("rel_volume"), 2)}x | Perf 1W={fmt(tv.get("perf_week"))}% | Perf 1M={fmt(tv.get("perf_month"))}%"""
    else:
        tv_text = "TEKNIKAL TRADINGVIEW: Tidak tersedia. Analisis hanya dari data BEI."

    return f"""Kamu analis saham BEI spesialis swing trading 1-5 hari. Tulis analisis kuantitatif dalam bahasa Indonesia.

SAHAM: {ticker} - {company}

DATA BEI:
- Close: {fmt(close)} | Change: {fmt(change)}%
- ATR-1d: {atr_pct:.2f}% = {atr_abs:.2f} absolut
- Net Foreign: {foreign_net:,.0f} (rasio: {foreign_ratio:.3f})
- Bid-Offer Pressure: {bid_pressure:+.3f} (>0 berarti buyer dominan)
- Signal Strength: {signal:.1f}/100 | Final Score: {final_score:.1f}/100
- Momentum: {safe_num(bei_row.get("momentum_score"), 50):.1f}
- Likuiditas: {safe_num(bei_row.get("liquidity_score"), 50):.1f}
- Flow: {safe_num(bei_row.get("flow_score"), 50):.1f}
- Market Activity: {safe_num(bei_row.get("market_activity_score"), 50):.1f}
- Volume Trend: {safe_num(bei_row.get("volume_trend_score"), 50):.1f}
- Price Structure: {safe_num(bei_row.get("price_structure_score"), 50):.1f}
{multiday_text}
{regime_text}
{tv_text}

LEVEL AWAL: Stop={fmt(close - atr_abs * 1.5)} | TP1={fmt(close + atr_abs * 2.0)} | TP2={fmt(close + atr_abs * 3.5)}

ATURAN:
- Jangan mengarang level harga.
- Jelaskan indikator yang bertentangan.
- Jika confidence rendah, gunakan WAIT.
- Bedakan data BEI, TradingView, dan asumsi.

FORMAT:
**Verdict:** [1 kalimat]

**Skor Keyakinan (0-100):**
- Total Confidence:
- Teknikal:
- Flow:
- Regime:
- Decision Tag: [AGRESIF BUY | BUY ON PULLBACK | WAIT | AVOID]

**Validasi Konsistensi Data:**
- Konsisten:
- Konflik:
- Kesimpulan:

**Kondisi Teknikal:**
RSI, MACD, ADX, EMA, Bollinger, volume relatif, dan performa.

**Kondisi Flow:**
Foreign flow, bid-offer pressure, broker activity, dan likuiditas.

**Peta Timeframe:**
- Intraday:
- Swing 1-3 hari:
- Swing 1-2 minggu:

**Skenario Eksekusi:**
1. Breakout continuation: trigger, entry, stop, TP, probabilitas.
2. Pullback setup: trigger, entry, stop, TP, probabilitas.
3. Failed breakout: invalidasi dan aksi.

**Risk Memo:**
Max risk per trade, volatility risk, liquidity trap risk, dan event risk.

**Rencana Final:**
Entry, stop, TP, R/R, holding plan, dan kapan review ulang.

**Red Flag:**
Jujur tentang kelemahan setup.

Target 450-650 kata."""


@st.cache_data(ttl=600, show_spinner=False)
def call_openrouter(prompt, api_key, model, max_tokens=2200):
    payload = {
        "model": model,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.35,
        "max_tokens": max_tokens,
        "top_p": 0.85,
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "https://sahambei.streamlit.app",
        "X-Title": "BEI Stock Screener",
    }
    try:
        response = requests.post(
            "https://openrouter.ai/api/v1/chat/completions",
            json=payload,
            headers=headers,
            timeout=40,
        )
        response.raise_for_status()
        body = response.json()
        message = ((body.get("choices") or [{}])[0].get("message") or {})
        return message.get("content") or "Model tidak mengembalikan konten teks."
    except requests.exceptions.HTTPError:
        code = response.status_code
        return {
            400: "Request invalid. Cek model/payload OpenRouter.",
            401: "API key OpenRouter tidak valid atau belum aktif.",
            402: "Kredit OpenRouter tidak cukup.",
            429: "Rate limit OpenRouter. Tunggu lalu coba lagi.",
        }.get(code, f"HTTP Error {code}")
    except Exception as exc:
        return f"Gagal: {exc}"


# ── LOCAL STOCK STORAGE ──────────────────────────────────────────────────────
def ensure_data_dir():
    os.makedirs(DATA_DIR, exist_ok=True)


def read_json_file(path, default):
    try:
        if not os.path.exists(path):
            return default
        with open(path, "r", encoding="utf-8") as handle:
            return json.load(handle)
    except Exception:
        return default


def write_json_file(path, payload):
    ensure_data_dir()
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)


def load_watchlist():
    items = read_json_file(WATCHLIST_FILE, [])
    return items if isinstance(items, list) else []


def save_watchlist(items):
    deduped = {}
    for item in items:
        if isinstance(item, dict) and item.get("id") and item.get("type") == "stock":
            deduped[item["id"]] = item
    write_json_file(WATCHLIST_FILE, list(deduped.values()))


def load_history():
    rows = read_json_file(HISTORY_FILE, [])
    return rows if isinstance(rows, list) else []


def save_history(rows):
    save_rows = [row for row in rows if isinstance(row, dict) and row.get("type") == "stock"]
    write_json_file(HISTORY_FILE, save_rows[-MAX_HISTORY_ROWS:])


def append_history_rows(rows):
    if rows:
        save_history(load_history() + rows)


def load_alert_rules():
    rules = DEFAULT_ALERT_RULES.copy()
    stored = read_json_file(ALERT_RULES_FILE, {})
    if isinstance(stored, dict):
        for key, value in stored.items():
            if key in rules:
                rules[key] = safe_num(value, rules[key])
    return rules


def save_alert_rules(rules):
    cleaned = {
        key: safe_num(rules.get(key), default)
        for key, default in DEFAULT_ALERT_RULES.items()
    }
    write_json_file(ALERT_RULES_FILE, cleaned)


def build_local_backup_payload():
    return {
        "schema": "stock_screener_backup_v1",
        "created_at": utc_now_iso(),
        "watchlist": load_watchlist(),
        "history": load_history(),
        "alert_rules": load_alert_rules(),
    }


def restore_local_backup_payload(payload, merge=True):
    if not isinstance(payload, dict) or payload.get("schema") not in {
        "stock_screener_backup_v1",
        "market_screener_backup_v1",
    }:
        raise ValueError("Format backup tidak dikenali.")
    watchlist = [item for item in payload.get("watchlist", []) if isinstance(item, dict) and item.get("type") == "stock"]
    history = [row for row in payload.get("history", []) if isinstance(row, dict) and row.get("type") == "stock"]
    alert_rules = payload.get("alert_rules", {})
    if not isinstance(alert_rules, dict):
        raise ValueError("Alert rules pada backup tidak valid.")

    if merge:
        current_watchlist = {item.get("id"): item for item in load_watchlist()}
        current_watchlist.update({item.get("id"): item for item in watchlist if item.get("id")})
        save_watchlist(list(current_watchlist.values()))

        current_history = load_history()
        known = {
            (row.get("ts"), row.get("id"), str(row.get("price")), str(row.get("score")))
            for row in current_history
        }
        for row in history:
            key = (row.get("ts"), row.get("id"), str(row.get("price")), str(row.get("score")))
            if key not in known:
                current_history.append(row)
                known.add(key)
        save_history(current_history)
        merged_rules = load_alert_rules()
        for key, value in alert_rules.items():
            if key in DEFAULT_ALERT_RULES:
                merged_rules[key] = safe_num(value, DEFAULT_ALERT_RULES[key])
        save_alert_rules(merged_rules)
    else:
        save_watchlist(watchlist)
        save_history(history)
        save_alert_rules(alert_rules)


def get_secret_source():
    if get_static_openrouter_key():
        try:
            if str(st.secrets.get("OPENROUTER_API_KEY", "")).strip():
                return "Streamlit secrets"
        except Exception:
            pass
        return ".env / environment"
    return "Belum diset"


def get_deploy_readiness_rows():
    secret_source = get_secret_source()
    try:
        ensure_data_dir()
        probe_path = os.path.join(DATA_DIR, ".write_probe")
        with open(probe_path, "w", encoding="utf-8") as handle:
            handle.write(utc_now_iso())
        os.remove(probe_path)
        storage_status = "OK"
        storage_note = "Folder data/ bisa ditulis."
    except Exception as exc:
        storage_status = "Error"
        storage_note = f"Folder data/ tidak bisa ditulis: {exc}"
    return pd.DataFrame([
        {
            "Area": "Secrets",
            "Status": "OK" if secret_source != "Belum diset" else "Perlu set",
            "Catatan": f"OpenRouter source: {secret_source}",
        },
        {
            "Area": "Config",
            "Status": "OK" if os.path.exists(".streamlit/config.toml") else "Kurang",
            "Catatan": ".streamlit/config.toml tersedia untuk deploy.",
        },
        {"Area": "Storage", "Status": storage_status, "Catatan": storage_note},
        {
            "Area": "Persistence",
            "Status": "Perhatian",
            "Catatan": "Cloud filesystem tidak permanen; gunakan backup JSON jika history penting.",
        },
        {
            "Area": "Dependencies",
            "Status": "OK",
            "Catatan": "requirements.txt dipakai saat build Streamlit.",
        },
    ])


def render_deploy_readiness_panel():
    with st.expander("Deploy Readiness", expanded=False):
        st.dataframe(get_deploy_readiness_rows(), width="stretch")
        st.caption("Untuk Streamlit Cloud, isi OPENROUTER_API_KEY dan OPENROUTER_MODEL di Secrets.")


def is_watchlisted(item_id):
    return any(item.get("id") == item_id for item in load_watchlist())


def add_watchlist_item(item):
    item = item.copy()
    item["added_at"] = item.get("added_at") or utc_now_iso()
    save_watchlist([old for old in load_watchlist() if old.get("id") != item.get("id")] + [item])


def remove_watchlist_item(item_id):
    save_watchlist([item for item in load_watchlist() if item.get("id") != item_id])


def watch_item_from_stock_row(row):
    ticker = str(row.get("kode saham", row.get("name", ""))).upper().strip()
    symbol = str(row.get("symbol") or f"IDX:{ticker}").upper().strip()
    company = str(row.get("nama perusahaan", row.get("description", ticker))).strip()
    return {
        "id": f"stock:{symbol}",
        "type": "stock",
        "source": "TradingView IDX",
        "symbol": symbol,
        "ticker": ticker,
        "label": ticker,
        "company": company,
        "sector": row.get("sector", "N/A"),
    }


# ── STOCK AUTO SCANNER ───────────────────────────────────────────────────────
def tv_scan_headers():
    return {
        "Content-Type": "application/json",
        "Origin": "https://www.tradingview.com",
        "Referer": "https://www.tradingview.com/",
        "User-Agent": "Mozilla/5.0",
    }


@st.cache_data(ttl=300, show_spinner=False)
def fetch_tv_stock_screener(limit=250, sort_by="Value.Traded", sort_order="desc"):
    payload = {
        "columns": STOCK_AUTO_FIELDS,
        "filter": [{"left": "type", "operation": "equal", "right": "stock"}],
        "markets": ["indonesia"],
        "options": {"lang": "id"},
        "range": [0, int(limit)],
        "sort": {"sortBy": sort_by, "sortOrder": sort_order},
        "symbols": {"query": {"types": []}, "tickers": []},
    }
    response = requests.post(
        "https://scanner.tradingview.com/indonesia/scan",
        json=payload,
        headers=tv_scan_headers(),
        timeout=25,
    )
    response.raise_for_status()
    body = response.json()
    return body.get("data", []), int(body.get("totalCount", 0) or 0)


@st.cache_data(ttl=300, show_spinner=False)
def fetch_tv_symbol_snapshot(symbol):
    payload = {
        "columns": STOCK_AUTO_FIELDS,
        "symbols": {"tickers": [symbol], "query": {"types": []}},
    }
    response = requests.post(
        "https://scanner.tradingview.com/indonesia/scan",
        json=payload,
        headers=tv_scan_headers(),
        timeout=15,
    )
    response.raise_for_status()
    data = response.json().get("data", [])
    if not data:
        return {}
    values = data[0].get("d", [])
    result = {field: values[index] for index, field in enumerate(STOCK_AUTO_FIELDS) if index < len(values)}
    result["symbol"] = data[0].get("s", symbol)
    return result


def rec_score(label):
    return {
        "Strong Buy": 100.0,
        "Buy": 78.0,
        "Neutral": 52.0,
        "Sell": 28.0,
        "Strong Sell": 8.0,
    }.get(str(label), 50.0)


def rsi_setup_score(value):
    rsi = safe_num(value, np.nan)
    if pd.isna(rsi):
        return 50.0
    if 45 <= rsi <= 65:
        return 90.0
    if 35 <= rsi < 45 or 65 < rsi <= 72:
        return 68.0
    if 30 <= rsi < 35 or 72 < rsi <= 80:
        return 48.0
    return 25.0


def bb_setup_score(value):
    position = safe_num(value, np.nan)
    if pd.isna(position):
        return 50.0
    if 45 <= position <= 85:
        return 86.0
    if 25 <= position < 45:
        return 64.0
    if 85 < position <= 100:
        return 55.0
    return 35.0


def normalize_stock_auto_rows(rows):
    records = []
    for item in rows:
        raw = item.get("d", [])
        tv = {
            field: raw[index] if index < len(raw) else None
            for index, field in enumerate(STOCK_AUTO_FIELDS)
        }
        symbol = str(item.get("s") or tv.get("symbol") or "")
        ticker = str(tv.get("name") or symbol.split(":")[-1]).upper().strip()
        if not ticker or ticker == "NAN":
            continue
        records.append({
            "kode saham": ticker,
            "symbol": symbol or f"IDX:{ticker}",
            "nama perusahaan": str(tv.get("description") or ticker).strip(),
            "sector": str(tv.get("sector") or "N/A").strip(),
            "close": safe_num(tv.get("close"), np.nan),
            "change_pct": safe_num(tv.get("change"), np.nan),
            "volume": safe_num(tv.get("volume"), np.nan),
            "value_traded": safe_num(tv.get("Value.Traded"), np.nan),
            "market_cap": safe_num(tv.get("market_cap_basic"), np.nan),
            **parse_tv_data(tv),
        })
    df = pd.DataFrame(records)
    if df.empty:
        return df
    for column in [
        "change_pct", "perf_week", "perf_month", "volume", "value_traded",
        "market_cap", "rel_volume", "adx",
    ]:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")
    df["rec_score"] = df["rec_all"].apply(rec_score)
    df["ema_score"] = df["ema_trend_count"].fillna(0).clip(0, 3) / 3 * 100
    df["adx_score"] = np.clip((df["adx"].fillna(15) - 12) * 4.5, 0, 100)
    df["macd_score"] = np.where(
        df["macd_cross"].eq("bullish"),
        78,
        np.where(df["macd_cross"].eq("bearish"), 34, 50),
    )
    df["trend_score"] = (
        df["rec_score"] * 0.35
        + df["ema_score"] * 0.25
        + df["adx_score"] * 0.20
        + df["macd_score"] * 0.20
    )
    df["momentum_score"] = (
        percentile_series(df["change_pct"].fillna(0)) * 0.40
        + percentile_series(df["perf_week"].fillna(0)) * 0.35
        + percentile_series(df["perf_month"].fillna(0)) * 0.25
    )
    df["liquidity_score"] = (
        percentile_series(df["value_traded"].fillna(0)) * 0.70
        + percentile_series(df["volume"].fillna(0)) * 0.30
    )
    df["volume_score"] = np.clip(df["rel_volume"].fillna(1.0) * 48, 0, 100)
    df["setup_score"] = (
        df["rsi"].apply(rsi_setup_score) * 0.55
        + df["bb_position"].apply(bb_setup_score) * 0.45
    )
    df["auto_score"] = (
        df["trend_score"] * 0.30
        + df["momentum_score"] * 0.25
        + df["liquidity_score"] * 0.20
        + df["volume_score"] * 0.15
        + df["setup_score"] * 0.10
    )

    def signal(row):
        score = safe_num(row.get("auto_score"), 0)
        change = safe_num(row.get("change_pct"), 0)
        relative_volume = safe_num(row.get("rel_volume"), 1)
        recommendation = rec_score(row.get("rec_all"))
        week = safe_num(row.get("perf_week"), 0)
        rsi = safe_num(row.get("rsi"), 50)
        if score >= 78 and change > 0 and relative_volume >= 1.15 and recommendation >= 70:
            return "Breakout Candidate"
        if score >= 68 and week > 0 and recommendation >= 52:
            return "Trending Up"
        if score >= 62 and change <= 0 and 38 <= rsi <= 58:
            return "Pullback Watch"
        if score >= 55:
            return "Watch Only"
        return "Weak / Avoid"

    df["signal_label"] = df.apply(signal, axis=1)
    quality_columns = [
        "close", "value_traded", "rel_volume", "rsi", "macd", "macd_signal",
        "ema20", "ema50", "ema200", "adx", "rec_all",
    ]
    df["data_quality_score"] = df.apply(
        lambda row: round(
            sum(
                1 for column in quality_columns
                if column in row.index
                and row.get(column) not in (None, "", "N/A")
                and not (isinstance(row.get(column), (float, np.floating)) and pd.isna(row.get(column)))
            ) / len(quality_columns) * 100,
            0,
        ),
        axis=1,
    )
    return df.sort_values("auto_score", ascending=False).reset_index(drop=True)


def fetch_auto_market_regime():
    raw = fetch_tv_symbol_snapshot("IDX:COMPOSITE")
    if not raw:
        return {}
    parsed = parse_tv_data(raw)
    change = safe_num(parsed.get("change_pct"), 0)
    recommendation = rec_score(parsed.get("rec_all"))
    adx = safe_num(parsed.get("adx"), 15)
    performance_week = safe_num(parsed.get("perf_week"), 0)
    score = float(np.clip(
        50 + change * 7
        + (recommendation - 50) * 0.35
        + min(adx, 35) * 0.35
        + performance_week * 1.2,
        0,
        100,
    ))
    return {
        "index_code": "COMPOSITE",
        "index_change_pct": change,
        "regime_score": score,
        "regime_label": "Risk-On" if score >= 62 else "Netral" if score >= 45 else "Risk-Off",
        "rec_all": parsed.get("rec_all", "N/A"),
        "rsi": parsed.get("rsi", np.nan),
        "adx": parsed.get("adx", np.nan),
        "perf_week": parsed.get("perf_week", np.nan),
        "perf_month": parsed.get("perf_month", np.nan),
    }


def auto_risk_flags(row, market_regime=None):
    """Return transparent warnings for a TradingView-only snapshot."""
    flags = []
    rsi = safe_num(row.get("rsi"), np.nan)
    rel_volume = safe_num(row.get("rel_volume"), np.nan)
    bb_position = safe_num(row.get("bb_position"), np.nan)
    change = safe_num(row.get("change_pct"), np.nan)
    adx = safe_num(row.get("adx"), np.nan)
    if not pd.isna(rsi) and rsi >= 75:
        flags.append("RSI tinggi: rawan pullback")
    elif not pd.isna(rsi) and rsi <= 30:
        flags.append("RSI rendah: momentum lemah atau perlu konfirmasi reversal")
    if not pd.isna(rel_volume) and rel_volume < 0.70:
        flags.append("Relative volume rendah")
    if not pd.isna(bb_position) and bb_position > 100:
        flags.append("Harga di atas Bollinger upper")
    if not pd.isna(change) and abs(change) >= 8:
        flags.append("Perubahan harian ekstrem")
    if not pd.isna(adx) and adx < 15:
        flags.append("ADX rendah: tren belum kuat")
    if market_regime and market_regime.get("regime_label") == "Risk-Off":
        flags.append("Market regime Risk-Off")
    if safe_num(row.get("data_quality_score"), 0) < 80:
        flags.append("Sebagian indikator TradingView tidak tersedia")
    return flags


def build_stock_auto_prompt(row, market_regime=None):
    close = safe_num(row.get("close"), np.nan)
    change = safe_num(row.get("change_pct"), np.nan)
    value_traded = safe_num(row.get("value_traded"), np.nan)
    relative_volume = safe_num(row.get("rel_volume"), np.nan)
    score = safe_num(row.get("auto_score"), 0)
    atr_proxy = max(abs(safe_num(row.get("change_pct"), 1.8)), 1.2)
    stop = close * (1 - atr_proxy * 1.35 / 100) if not pd.isna(close) else np.nan
    tp1 = close * (1 + atr_proxy * 1.8 / 100) if not pd.isna(close) else np.nan
    tp2 = close * (1 + atr_proxy * 3.0 / 100) if not pd.isna(close) else np.nan
    regime_text = "Market regime otomatis tidak tersedia."
    if market_regime:
        regime_text = (
            f"Market regime: {market_regime.get('regime_label')} "
            f"(score={safe_num(market_regime.get('regime_score'), 0):.1f}, "
            f"COMPOSITE={pct_text(market_regime.get('index_change_pct'))}, "
            f"rec={market_regime.get('rec_all', 'N/A')})"
        )
    return f"""Kamu analis saham BEI untuk swing trading. Data berasal dari TradingView scanner otomatis.

SAHAM: {row.get('kode saham')} - {row.get('nama perusahaan')}
SEKTOR: {row.get('sector')}

DATA:
- Close: {format_idr(close, compact=False)}
- Change harian: {pct_text(change)}
- Volume: {format_compact(row.get('volume'))}
- Value traded: {format_idr(value_traded)}
- Relative volume 10D: {ratio_text(relative_volume)}
- Market cap: {format_idr(row.get('market_cap'))}
- Rekomendasi: {row.get('rec_all')} | MA: {row.get('rec_ma')} | Oscillator: {row.get('rec_other')}
- RSI: {safe_num(row.get('rsi'), np.nan):.1f} | MACD: {row.get('macd_cross')}
- ADX: {safe_num(row.get('adx'), np.nan):.1f}
- EMA20/50/200: {format_idr(row.get('ema20'), compact=False)} / {format_idr(row.get('ema50'), compact=False)} / {format_idr(row.get('ema200'), compact=False)}
- Bollinger position: {safe_num(row.get('bb_position'), np.nan):.1f}%
- Performa 1W/1M: {pct_text(row.get('perf_week'))} / {pct_text(row.get('perf_month'))}

AUTO SCORE:
- Total: {score:.1f}/100
- Trend: {safe_num(row.get('trend_score'), 0):.1f}
- Momentum: {safe_num(row.get('momentum_score'), 0):.1f}
- Likuiditas: {safe_num(row.get('liquidity_score'), 0):.1f}
- Volume: {safe_num(row.get('volume_score'), 0):.1f}
- Setup: {safe_num(row.get('setup_score'), 0):.1f}
- Signal: {row.get('signal_label')}

REGIME:
{regime_text}

LEVEL PROXY:
- Stop awal: {format_idr(stop, compact=False)}
- TP1: {format_idr(tp1, compact=False)}
- TP2: {format_idr(tp2, compact=False)}

BATASAN:
- Jangan mengklaim foreign flow, broker summary, atau orderbook karena tidak tersedia di auto mode.
- Jika perlu konfirmasi, gunakan WAIT atau BUY ON CONFIRMATION.
- Semua level harus konsisten dengan data.

FORMAT:
**Verdict:**
**Decision Tag:** [AGRESIF BUY | BUY ON CONFIRMATION | BUY ON PULLBACK | WAIT | AVOID]
**Confidence:** [0-100] + alasan
**Analisis Teknikal:**
**Analisis Likuiditas:**
**Skenario 1-5 Hari:**
1. Continuation - trigger, entry, stop, TP.
2. Pullback - trigger, entry, stop, TP.
3. Invalidasi - trigger dan aksi.
**Yang Tidak Terlihat di Auto Mode:**
Jelaskan bahwa foreign flow dan broker summary tidak tersedia.
**Rencana Final:**
Entry, stop, TP, dan kapan review ulang."""


def tradingview_advanced_chart_url(symbol, interval="D"):
    params = {
        "frameElementId": f"tradingview_{str(symbol).replace(':', '_')}",
        "symbol": symbol,
        "interval": interval,
        "hidesidetoolbar": "0",
        "hide_top_toolbar": "0",
        "hide_legend": "0",
        "hide_volume": "0",
        "symboledit": "1",
        "saveimage": "1",
        "toolbarbg": "f1f3f6",
        "studies": "[]",
        "theme": "dark",
        "style": "1",
        "timezone": "Asia/Jakarta",
        "withdateranges": "1",
        "hideideas": "1",
        "locale": "id",
    }
    return "https://s.tradingview.com/widgetembed/?" + urlencode(params)


def render_tradingview_advanced_chart(symbol, interval="D", height=720):
    st.markdown(f"**TradingView Chart — {symbol} — timeframe {interval}**")
    components.iframe(
        tradingview_advanced_chart_url(symbol, interval),
        height=height,
        scrolling=False,
    )


def render_openrouter_sidebar():
    st.header("Setup")
    st.markdown("**OpenRouter API Key**")
    static_key = get_static_openrouter_key()
    model = get_openrouter_model()
    st.caption(f"Model aktif: {model}")
    if static_key:
        st.success("API key statis terdeteksi.")
        with st.expander("Override API Key (opsional)"):
            override = st.text_input("Override Key", type="password", key="or_key_override")
        key = override.strip() if override else static_key
    else:
        key = st.text_input("API Key", type="password", placeholder="sk-or-v1-...", key="or_key")
    return key.strip(), model


def render_stock_auto_page(openrouter_key, llm_model):
    st.title("BEI Auto Scanner")
    st.caption("Saham Indonesia otomatis dari TradingView scanner. Tidak perlu upload file Excel.")
    goapi_configured = goapi_is_configured()
    with st.sidebar:
        st.header("Filter Auto Saham")
        if goapi_configured:
            st.success("GOAPI IDX aktif untuk quote + historical")
        else:
            st.info("GOAPI belum diatur; quote tetap dari TradingView")
        if st.button("Refresh data TradingView", width="stretch", key="stock_auto_refresh"):
            fetch_tv_stock_screener.clear()
            fetch_tv_symbol_snapshot.clear()
            fetch_tv_data.clear()
            clear_goapi_cache()
            st.session_state["stock_auto_last_refresh"] = utc_now_iso()
            st.rerun()
        chart_interval_labels = {
            "5 menit": "5",
            "15 menit": "15",
            "1 jam": "60",
            "4 jam": "240",
            "Harian": "D",
            "Mingguan": "W",
        }
        chart_interval_label = st.selectbox(
            "Timeframe chart",
            list(chart_interval_labels),
            index=4,
            key="stock_auto_chart_interval",
        )
        sort_label = st.selectbox("Urutkan", list(STOCK_SORT_OPTIONS.keys()), index=0, key="stock_auto_sort")
        order_label = st.selectbox("Arah", ["Desc", "Asc"], index=0, key="stock_auto_sort_order")
        fetch_limit = st.number_input("Ambil data", 50, 800, 250, 50, key="stock_auto_fetch_limit")
        min_value = st.number_input("Min value traded", 0, 5_000_000_000, 5_000_000_000, 1_000_000_000, key="stock_auto_min_value")
        min_score = st.slider("Min auto score", 0, 100, 55, key="stock_auto_min_score")
        top_n = st.number_input("Top N", 5, 200, 30, 5, key="stock_auto_top_n")
        search = st.text_input("Cari saham", placeholder="BBCA, BBRI, TLKM", key="stock_auto_search")

    try:
        with st.spinner("Mengambil saham otomatis dari TradingView..."):
            rows, total_count = fetch_tv_stock_screener(
                limit=fetch_limit,
                sort_by=STOCK_SORT_OPTIONS.get(sort_label, "Value.Traded"),
                sort_order="asc" if order_label == "Asc" else "desc",
            )
            df = normalize_stock_auto_rows(rows)
            market_regime = fetch_auto_market_regime()
    except Exception as exc:
        st.error(f"Auto scanner gagal mengambil data TradingView: {exc}")
        st.stop()

    if df.empty:
        st.warning("TradingView tidak mengembalikan data saham.")
        st.stop()

    refresh_text = st.session_state.get("stock_auto_last_refresh", "Cache maksimal 5 menit")
    st.caption(f"Sumber: TradingView Indonesia scanner · Refresh: {refresh_text}")

    with st.sidebar:
        sectors = sorted([item for item in df["sector"].dropna().astype(str).unique() if item and item != "N/A"])
        sector_filter = st.selectbox("Sektor", ["Semua"] + sectors, key="stock_auto_sector")
        signal_filter = st.selectbox(
            "Sinyal",
            ["Semua", "Breakout Candidate", "Trending Up", "Pullback Watch", "Watch Only", "Weak / Avoid"],
            key="stock_auto_signal",
        )

    view = df[df["value_traded"].fillna(0) >= float(min_value)].copy()
    view = view[view["auto_score"].fillna(0) >= float(min_score)]
    if sector_filter != "Semua":
        view = view[view["sector"] == sector_filter]
    if signal_filter != "Semua":
        view = view[view["signal_label"] == signal_filter]
    if search:
        query = search.upper().strip()
        view = view[
            view["kode saham"].astype(str).str.contains(query, na=False)
            | view["nama perusahaan"].astype(str).str.upper().str.contains(query, na=False)
        ]
    view = view.head(int(top_n))

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Universe", f"{total_count} saham")
    m2.metric("Ditampilkan", len(view))
    m3.metric("Top Value", format_idr(df["value_traded"].max()))
    m4.metric("Regime", market_regime.get("regime_label", "N/A") if market_regime else "N/A",
              pct_text(market_regime.get("index_change_pct")) if market_regime else None)

    if market_regime:
        regime_cols = st.columns(4)
        regime_cols[0].metric("IHSG change", pct_text(market_regime.get("index_change_pct")))
        regime_cols[1].metric("IHSG RSI", f"{safe_num(market_regime.get('rsi'), 0):.1f}")
        regime_cols[2].metric("IHSG ADX", f"{safe_num(market_regime.get('adx'), 0):.1f}")
        regime_cols[3].metric("IHSG 1W", pct_text(market_regime.get("perf_week")))

    st.subheader("Kandidat Saham Auto")
    table_cols = [
        "kode saham", "nama perusahaan", "sector", "close", "change_pct",
        "value_traded", "rel_volume", "rec_all", "rsi", "adx",
        "perf_week", "perf_month", "auto_score", "data_quality_score", "signal_label",
    ]
    display = view[[column for column in table_cols if column in view.columns]].copy()
    if not display.empty:
        display["close"] = display["close"].apply(lambda value: format_idr(value, compact=False))
        display["value_traded"] = display["value_traded"].apply(format_idr)
        for column in ["change_pct", "perf_week", "perf_month"]:
            if column in display.columns:
                display[column] = display[column].apply(pct_text)
        display["rel_volume"] = display["rel_volume"].apply(ratio_text)
        render_df_with_style_fallback(display, ["auto_score"])
    else:
        st.info("Tidak ada saham sesuai filter.")
    st.download_button(
        "Export Excel",
        dataframe_to_excel_bytes({"Kandidat Auto": view}),
        "bei_auto_tradingview.xlsx",
        EXCEL_MIME,
        help="Workbook Excel berisi semua row dan kolom hasil scanner; header sudah freeze dan filter aktif.",
    )

    if view.empty:
        st.warning("Tidak ada kandidat setelah filter.")
        st.stop()

    st.markdown("---")
    selected_code = st.selectbox("Pilih saham untuk analisis", view["kode saham"].tolist(), key="stock_auto_selected")
    selected = view[view["kode saham"] == selected_code].iloc[0]
    st.subheader(f"{selected['kode saham']}  {selected['nama perusahaan']}  {selected['signal_label']}")
    m1, m2, m3, m4, m5, m6 = st.columns(6)
    m1.metric("Auto Score", f"{safe_num(selected.get('auto_score'), 0):.1f}")
    m2.metric("Close", format_idr(selected.get("close"), compact=False))
    m3.metric("Change", pct_text(selected.get("change_pct")))
    m4.metric("Value", format_idr(selected.get("value_traded")))
    m5.metric("TV Rec", str(selected.get("rec_all", "N/A")))
    m6.metric("Data quality", f"{safe_num(selected.get('data_quality_score'), 0):.0f}%")

    goapi_snapshot = {}
    goapi_history = pd.DataFrame()
    goapi_broker_rows = []
    goapi_quant = {
        "observations": 0,
        "status": "insufficient_data",
        "indicators": {},
        "levels": {"support": [], "resistance": [], "fibonacci": []},
        "execution": {},
        "history": [],
    }
    if goapi_configured:
        try:
            with st.spinner(f"Mengambil historical {selected_code} dari GOAPI..."):
                goapi_snapshot = fetch_goapi_snapshot(selected_code)
                goapi_history = fetch_goapi_historical(selected_code)
                if not goapi_history.empty:
                    goapi_quant = build_quant_context(goapi_history, goapi_snapshot)
                    latest_date = goapi_history["date"].dropna().max()
                    if not pd.isna(latest_date):
                        try:
                            goapi_broker_rows = fetch_goapi_broker_summary(
                                selected_code,
                                pd.Timestamp(latest_date).strftime("%Y-%m-%d"),
                                "ALL",
                            )
                        except GoAPIError:
                            # A trial key may expose historical data but not broker summary.
                            goapi_broker_rows = []
            if goapi_snapshot or not goapi_history.empty:
                st.caption(
                    f"Backend data: GOAPI IDX · {len(goapi_history)} candle historis · "
                    f"quote {format_idr(goapi_snapshot.get('close'), compact=False)} "
                    f"({pct_text(goapi_snapshot.get('change_pct'))})"
                )
        except GoAPIError as exc:
            st.warning(f"GOAPI tidak tersedia untuk {selected_code}: {exc}. Data scanner tetap memakai TradingView.")

    risk_flags = auto_risk_flags(selected, market_regime)
    if risk_flags:
        st.warning(" | ".join(risk_flags))

    stock_item = watch_item_from_stock_row(selected)
    a1, a2, a3 = st.columns([1, 1, 2])
    with a1:
        if is_watchlisted(stock_item["id"]):
            if st.button("Hapus Watchlist", width="stretch", key=f"stock_rm_{stock_item['id']}"):
                remove_watchlist_item(stock_item["id"])
                st.rerun()
        elif st.button("Tambah Watchlist", width="stretch", key=f"stock_add_{stock_item['id']}"):
            add_watchlist_item(stock_item)
            st.success("Saham masuk watchlist.")
    with a2:
        if st.button("Simpan Snapshot", width="stretch", key=f"stock_snap_{stock_item['id']}"):
            append_history_rows([history_row_from_stock(selected, stock_item["id"])])
            st.success("Snapshot saham tersimpan.")
    with a3:
        st.caption("Watchlist dan history saham tersimpan lokal.")

    section = st.radio(
        "Detail Saham Auto",
        ["Score", "Teknikal", "Chart", "News", "History", "AI Analisis"],
        horizontal=True,
        index=2,
        label_visibility="collapsed",
        key=f"stock_auto_detail_{selected_code}",
    )
    if section == "Score":
        breakdown = pd.DataFrame({
            "Faktor": ["Trend", "Momentum", "Likuiditas", "Volume Relatif", "Setup"],
            "Skor": [
                safe_num(selected.get("trend_score"), np.nan),
                safe_num(selected.get("momentum_score"), np.nan),
                safe_num(selected.get("liquidity_score"), np.nan),
                safe_num(selected.get("volume_score"), np.nan),
                safe_num(selected.get("setup_score"), np.nan),
            ],
        })
        breakdown["Status"] = breakdown["Skor"].apply(factor_label)
        render_df_with_style_fallback(breakdown, ["Skor"])
        st.info("Mode Auto tidak membaca net foreign dan broker summary. Gunakan Upload BEI Advanced untuk flow detail.")
    elif section == "Teknikal":
        ta1, ta2, ta3 = st.columns(3)
        ta1.metric("Rekomendasi", str(selected.get("rec_all", "N/A")))
        ta2.metric("MA Signal", str(selected.get("rec_ma", "N/A")))
        ta3.metric("Oscillator", str(selected.get("rec_other", "N/A")))
        tb1, tb2, tb3, tb4 = st.columns(4)
        tb1.metric("RSI(14)", f"{safe_num(selected.get('rsi'), 0):.1f}")
        tb2.metric("ADX", f"{safe_num(selected.get('adx'), 0):.1f}")
        tb3.metric("Stoch K/D", f"{safe_num(selected.get('stoch_k'), 0):.1f}/{safe_num(selected.get('stoch_d'), 0):.1f}")
        tb4.metric("Vol Rel", ratio_text(selected.get("rel_volume")))
        tc1, tc2, tc3 = st.columns(3)
        for column, label, key in [(tc1, "EMA20", "above_ema20"), (tc2, "EMA50", "above_ema50"), (tc3, "EMA200", "above_ema200")]:
            state = selected.get(key)
            column.metric(label, format_idr(selected.get(label.lower()), compact=False), delta="di atas" if state else "di bawah" if state is False else "-")
        if selected.get("bb_position") is not None and not pd.isna(selected.get("bb_position")):
            st.progress(min(max(int(safe_num(selected.get("bb_position"), 0)), 0), 100),
                        text=f"Bollinger Position: {safe_num(selected.get('bb_position'), 0):.1f}%")
        if not goapi_history.empty:
            indicators = goapi_quant.get("indicators", {})
            st.markdown("**Historical OHLC dari GOAPI**")
            go1, go2, go3, go4 = st.columns(4)
            go1.metric("Candle", len(goapi_history))
            go2.metric("RSI14", price_text(indicators.get("rsi14")))
            go3.metric("ATR14", price_text(indicators.get("atr14")))
            go4.metric("Status", goapi_quant.get("status", "N/A"))
            st.caption(
                f"Support: {goapi_quant.get('levels', {}).get('support', [])} · "
                f"Resistance: {goapi_quant.get('levels', {}).get('resistance', [])} · "
                f"Broker summary: {len(goapi_broker_rows)} baris"
            )
    elif section == "Chart":
        render_tradingview_advanced_chart(
            f"IDX:{selected_code}",
            interval=chart_interval_labels[chart_interval_label],
            height=760,
        )
        st.caption("Gunakan toolbar kiri untuk trendline, horizontal line, Fibonacci, dan drawing lain. Data flow foreign/broker tidak tersedia di Auto TradingView.")
        st.markdown(
            f"[Buka chart penuh di TradingView](https://www.tradingview.com/chart/?symbol=IDX%3A{selected_code})"
        )
    elif section == "News":
        components.iframe(
            f"https://id.tradingview.com/symbols/IDX-{selected_code}/news/",
            height=720,
            scrolling=True,
        )
    elif section == "History":
        history = pd.DataFrame(load_history())
        item_history = history[history["id"] == stock_item["id"]].copy() if not history.empty and "id" in history.columns else pd.DataFrame()
        if item_history.empty:
            st.info("Belum ada history. Klik Simpan Snapshot atau refresh Watchlist.")
        else:
            item_history["ts"] = pd.to_datetime(item_history["ts"], errors="coerce")
            metrics = [column for column in ["score", "change", "rel_volume"] if column in item_history.columns]
            if metrics:
                st.line_chart(item_history.sort_values("ts").set_index("ts")[metrics])
            st.dataframe(item_history.sort_values("ts", ascending=False), width="stretch")
    elif section == "AI Analisis":
        analysis_rendered = False
        if not openrouter_key:
            st.warning("Masukkan OpenRouter API Key di sidebar untuk AI.")
        elif st.button("Generate Analisis AI", width="stretch", key=f"stock_auto_ai_{selected_code}"):
            with st.spinner("OpenRouter menganalisis saham otomatis..."):
                auto_context = build_runtime_context(
                    selected_code,
                    str(selected.get("nama perusahaan", selected_code)),
                    {},
                    {
                        **_ai_row_payload(selected),
                        "goapi_snapshot": goapi_snapshot,
                        "goapi_broker_summary": goapi_broker_rows,
                    },
                    goapi_quant,
                    market_regime,
                    {
                        "mode": "auto_tradingview",
                        "goapi_available": bool(not goapi_history.empty),
                        "goapi_broker_available": bool(goapi_broker_rows),
                        "note": (
                            "Historical OHLC dan broker summary berasal dari GOAPI; order book BEI tidak tersedia."
                            if goapi_history is not None and not goapi_history.empty and goapi_broker_rows else
                            "Historical OHLC berasal dari GOAPI; broker summary atau order book BEI tidak tersedia."
                            if not goapi_history.empty else
                            "OHLC historis GOAPI tidak tersedia; foreign flow, broker summary, dan order book BEI juga tidak tersedia."
                        ),
                    },
                )
                outcome = call_openrouter_structured(auto_context, openrouter_key, llm_model)
                if not outcome.get("ok"):
                    st.error(outcome.get("error", "Analisis AI gagal."))
                else:
                    try:
                        validated, warnings = validate_analysis(outcome["data"], selected_code, goapi_history, auto_context["quant_context"])
                        st.session_state[_ai_state_key(selected_code)] = validated
                        save_status = save_analysis(selected_code, auto_context, validated)
                        render_structured_analysis(
                            validated,
                            warnings,
                            chart_history=goapi_history,
                            quant_context=goapi_quant,
                        )
                        st.caption(f"Storage: lokal={save_status['local']} | Supabase={save_status['supabase']}")
                        analysis_rendered = True
                    except ValueError as exc:
                        st.error(f"Validasi respons AI gagal: {exc}")
        stored_analysis = st.session_state.get(_ai_state_key(selected_code))
        if stored_analysis and not analysis_rendered:
            st.caption("Hasil analisis terakhir pada sesi ini.")
            render_structured_analysis(
                stored_analysis,
                stored_analysis.get("validation_warnings", []),
                chart_history=goapi_history,
                quant_context=goapi_quant,
            )


# ── WATCHLIST AND ALERTS ─────────────────────────────────────────────────────
def history_row_from_stock(row, item_id=None):
    ticker = str(row.get("kode saham", row.get("name", ""))).upper()
    symbol = str(row.get("symbol") or f"IDX:{ticker}").upper()
    return {
        "ts": utc_now_iso(),
        "id": item_id or f"stock:{symbol}",
        "type": "stock",
        "label": ticker,
        "price": safe_num(row.get("close"), np.nan),
        "score": safe_num(row.get("auto_score"), np.nan),
        "volume": safe_num(row.get("volume"), np.nan),
        "value_traded": safe_num(row.get("value_traded"), np.nan),
        "change": safe_num(row.get("change_pct"), np.nan),
        "rel_volume": safe_num(row.get("rel_volume"), np.nan),
        "status": row.get("signal_label", ""),
        "rec": row.get("rec_all", ""),
    }


def evaluate_alerts_for_stock(row, rules):
    alerts = []
    label = str(row.get("kode saham", row.get("name", ""))).upper()
    score = safe_num(row.get("auto_score"), 0)
    change = safe_num(row.get("change_pct"), 0)
    relative_volume = safe_num(row.get("rel_volume"), 0)
    signal = str(row.get("signal_label", ""))
    if score >= rules["stock_min_score"]:
        alerts.append({"Level": "Info", "Asset": label, "Alert": f"Auto score >= {rules['stock_min_score']:.0f}", "Value": f"{score:.1f}"})
    if change >= rules["stock_min_change"]:
        alerts.append({"Level": "Momentum", "Asset": label, "Alert": f"Change >= {rules['stock_min_change']:.1f}%", "Value": pct_text(change)})
    if change <= rules["stock_dump_change"]:
        alerts.append({"Level": "Risk", "Asset": label, "Alert": f"Drop <= {rules['stock_dump_change']:.1f}%", "Value": pct_text(change)})
    if relative_volume >= rules["stock_min_rel_volume"]:
        alerts.append({"Level": "Volume", "Asset": label, "Alert": f"Rel volume >= {rules['stock_min_rel_volume']:.1f}x", "Value": ratio_text(relative_volume)})
    if signal == "Breakout Candidate":
        alerts.append({"Level": "Momentum", "Asset": label, "Alert": "Breakout Candidate", "Value": f"{score:.1f}"})
    return alerts


def refresh_watchlist_snapshot():
    items = load_watchlist()
    snapshot_rows, history_rows, alerts = [], [], []
    rules = load_alert_rules()
    stock_items = [item for item in items if item.get("type") == "stock"]
    if not stock_items:
        return pd.DataFrame(), pd.DataFrame(), 0

    try:
        rows, _ = fetch_tv_stock_screener(limit=900, sort_by="Value.Traded", sort_order="desc")
        stock_df = normalize_stock_auto_rows(rows)
        for item in stock_items:
            symbol = str(item.get("symbol") or f"IDX:{item.get('ticker', '')}").upper()
            ticker = symbol.split(":")[-1]
            match = stock_df[stock_df["kode saham"] == ticker] if not stock_df.empty else pd.DataFrame()
            if match.empty:
                snapshot_rows.append({"id": item.get("id"), "type": "stock", "label": item.get("label"), "status": "Tidak ditemukan"})
                continue
            row = match.iloc[0]
            history_rows.append(history_row_from_stock(row, item.get("id")))
            alerts.extend(evaluate_alerts_for_stock(row, rules))
            snapshot_rows.append({
                "id": item.get("id"),
                "type": "stock",
                "label": row.get("kode saham"),
                "source": "TradingView IDX",
                "price": safe_num(row.get("close"), np.nan),
                "score": safe_num(row.get("auto_score"), np.nan),
                "volume": safe_num(row.get("volume"), np.nan),
                "value_traded": safe_num(row.get("value_traded"), np.nan),
                "change": safe_num(row.get("change_pct"), np.nan),
                "rel_volume": safe_num(row.get("rel_volume"), np.nan),
                "status": row.get("signal_label", ""),
            })
    except Exception as exc:
        snapshot_rows.append({"id": "stock:error", "type": "stock", "label": "TradingView IDX", "status": f"Error: {exc}"})

    append_history_rows(history_rows)
    return pd.DataFrame(snapshot_rows), pd.DataFrame(alerts), len(history_rows)


def render_home():
    st.title("BEI Stock Screener")
    st.caption("Analisis saham BEI dengan data TradingView, upload BEI, AI opsional, watchlist, dan alert.")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.subheader("Auto TradingView")
        st.write("Scan saham IDX tanpa upload file, lengkap dengan score teknikal dan kandidat setup.")
        if st.button("Buka Auto Scanner", width="stretch"):
            set_page("Saham BEI")
    with c2:
        st.subheader("Upload BEI Advanced")
        st.write("Gunakan ringkasan saham, broker, perdagangan, dan indeks untuk flow yang lebih detail.")
        if st.button("Buka Saham BEI", width="stretch"):
            set_page("Saham BEI")
    with c3:
        st.subheader("Watchlist & Alerts")
        st.write("Simpan saham, refresh snapshot, lihat history score, dan evaluasi alert.")
        if st.button("Buka Watchlist", width="stretch"):
            set_page("Watchlist & Alerts")
    st.markdown("---")
    st.info("Semua output adalah alat bantu analisis, bukan rekomendasi investasi.")
    render_deploy_readiness_panel()


def render_watchlist_page():
    st.title("Watchlist & Alerts Saham")
    st.caption("Watchlist, history score, dan alert saham IDX disimpan lokal di folder data/.")
    items = load_watchlist()

    with st.sidebar:
        st.header("Alert Rules Saham")
        rules = load_alert_rules()
        with st.form("stock_alert_rules_form"):
            rules["stock_min_score"] = st.number_input("Min auto score", 0.0, 100.0, float(rules["stock_min_score"]), 1.0)
            rules["stock_min_change"] = st.number_input("Naik harian >=", -100.0, 1000.0, float(rules["stock_min_change"]), 1.0)
            rules["stock_dump_change"] = st.number_input("Turun harian <=", -100.0, 0.0, float(rules["stock_dump_change"]), 1.0)
            rules["stock_min_rel_volume"] = st.number_input("Relative volume >=", 0.0, 20.0, float(rules["stock_min_rel_volume"]), 0.1)
            if st.form_submit_button("Simpan Rules", width="stretch"):
                save_alert_rules(rules)
                st.success("Alert rules tersimpan.")

    render_deploy_readiness_panel()
    with st.expander("Backup / Restore Data Saham", expanded=False):
        backup = json.dumps(build_local_backup_payload(), ensure_ascii=False, indent=2).encode("utf-8")
        st.download_button(
            "Export Backup JSON",
            backup,
            file_name=f"stock_screener_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json",
            mime="application/json",
            width="stretch",
        )
        uploaded = st.file_uploader("Restore backup JSON", type=["json"], key="restore_stock_backup")
        merge = st.checkbox("Merge dengan data sekarang", value=True)
        if uploaded is not None and st.button("Restore Backup", width="stretch"):
            try:
                restore_local_backup_payload(json.loads(uploaded.read().decode("utf-8")), merge=merge)
                st.success("Backup saham berhasil direstore.")
                st.rerun()
            except Exception as exc:
                st.error(f"Restore gagal: {exc}")

    if not items:
        st.info("Watchlist masih kosong. Tambahkan saham dari Auto Scanner.")
        return

    st.subheader("Daftar Watchlist")
    watch_df = pd.DataFrame(items)
    columns = [column for column in ["label", "company", "sector", "source", "symbol", "added_at"] if column in watch_df.columns]
    st.dataframe(watch_df[columns], width="stretch")

    a1, a2 = st.columns([1, 2])
    with a1:
        if st.button("Refresh Watchlist", width="stretch"):
            with st.spinner("Mengambil snapshot terbaru saham..."):
                snapshot, alerts, saved_count = refresh_watchlist_snapshot()
            st.session_state["stock_watchlist_snapshot"] = snapshot
            st.session_state["stock_watchlist_alerts"] = alerts
            st.success(f"{saved_count} snapshot tersimpan.")
    with a2:
        remove_options = {
            f"{item.get('label', item.get('id'))}": item.get("id")
            for item in items
        }
        selected_remove = st.selectbox("Hapus item", ["-"] + list(remove_options.keys()))
        if selected_remove != "-" and st.button("Hapus dari Watchlist"):
            remove_watchlist_item(remove_options[selected_remove])
            st.rerun()

    snapshot = st.session_state.get("stock_watchlist_snapshot")
    alerts = st.session_state.get("stock_watchlist_alerts")
    if isinstance(snapshot, pd.DataFrame) and not snapshot.empty:
        st.subheader("Snapshot Terbaru")
        st.dataframe(snapshot, width="stretch")
    else:
        st.info("Klik Refresh Watchlist untuk mengambil data terbaru dan menyimpan history.")

    st.subheader("Alert Aktif")
    if isinstance(alerts, pd.DataFrame) and not alerts.empty:
        st.dataframe(alerts, width="stretch")
    else:
        st.caption("Belum ada alert aktif dari snapshot terakhir.")

    st.subheader("History")
    history = pd.DataFrame(load_history())
    if history.empty:
        st.info("History masih kosong.")
        return
    history["ts"] = pd.to_datetime(history["ts"], errors="coerce")
    labels = {item.get("id"): item.get("label", item.get("id")) for item in items}
    ids = history["id"].dropna().unique().tolist()
    selected_label = st.selectbox("Asset history", [labels.get(item_id, item_id) for item_id in ids])
    selected_id = next(item_id for item_id in ids if labels.get(item_id, item_id) == selected_label)
    selected_history = history[history["id"] == selected_id].sort_values("ts")
    metrics = [column for column in ["score", "change", "rel_volume"] if column in selected_history.columns]
    if metrics:
        st.line_chart(selected_history.set_index("ts")[metrics])
    st.dataframe(selected_history.sort_values("ts", ascending=False), width="stretch")


def set_page(page):
    st.session_state["page"] = page
    st.rerun()


# ── UPLOAD AND ADVANCED SCORING ───────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def load_table(uploaded_file):
    content = io.BytesIO(uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read())
    if uploaded_file.name.lower().endswith((".xlsx", ".xls")):
        try:
            return normalize_columns(pd.read_excel(content, sheet_name=0))
        except ImportError as exc:
            raise RuntimeError(
                "Dependency Excel belum terpasang pada Python yang menjalankan Streamlit. "
                "Jalankan: venv/bin/python -m pip install -r requirements.txt"
            ) from exc
    return normalize_columns(pd.read_csv(content, sep=None, engine="python"))


def uploaded_file_payload(uploaded_file):
    if not uploaded_file:
        return {}
    return {
        "name": uploaded_file.name,
        "bytes": uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read(),
        "mime": getattr(uploaded_file, "type", None) or "application/octet-stream",
    }


def compute_scores(saham_df, broker_df, trade_df, daftar_df, weights=None):
    weights = weights or DEFAULT_WEIGHTS
    data = saham_df.copy()
    data["kode saham"] = data[find_col(data, "kode saham")].astype(str).str.upper().str.strip()
    data["nama perusahaan"] = data[find_col(data, "nama perusahaan")].astype(str).str.strip()
    # Keep the official trading date with the scored row so the AI chart and
    # future persistence use the actual market date, not the upload label.
    try:
        data["trading_date"] = data[find_col(data, "tanggal perdagangan terakhir")]
    except KeyError:
        data["trading_date"] = pd.NaT
    for column, alias in [
        ("penutupan", "penutupan_num"), ("open price", "open_num"),
        ("tertinggi", "high_num"), ("terendah", "low_num"),
        ("sebelumnya", "prev_num"), ("selisih", "selisih_num"),
        ("nilai", "nilai_num"), ("volume", "volume_num"),
        ("frekuensi", "frekuensi_num"), ("foreign buy", "foreign_buy_num"),
        ("foreign sell", "foreign_sell_num"), ("bid volume", "bid_vol_num"),
        ("offer volume", "offer_vol_num"),
    ]:
        data[alias] = safe_col(data, column)
    data["change_pct"] = np.where(data["prev_num"] != 0, data["selisih_num"] / data["prev_num"] * 100, np.nan)
    data["foreign_net"] = data["foreign_buy_num"] - data["foreign_sell_num"]
    data["foreign_net_ratio"] = np.where(data["nilai_num"] != 0, data["foreign_net"] / data["nilai_num"], 0)
    data["bid_offer_pressure"] = np.where(
        data["bid_vol_num"] + data["offer_vol_num"] != 0,
        (data["bid_vol_num"] - data["offer_vol_num"]) / (data["bid_vol_num"] + data["offer_vol_num"]),
        0,
    )
    data["true_range"] = np.maximum(
        data["high_num"] - data["low_num"],
        np.maximum(abs(data["high_num"] - data["prev_num"]), abs(data["low_num"] - data["prev_num"])),
    )
    data["atr_pct"] = np.where(data["prev_num"] != 0, data["true_range"] / data["prev_num"] * 100, np.nan)
    data["close_position"] = np.where(
        data["high_num"] - data["low_num"] != 0,
        (data["penutupan_num"] - data["low_num"]) / (data["high_num"] - data["low_num"]),
        0.5,
    )

    listing = daftar_df.copy()
    if has_columns(listing, REQUIRED_COLUMNS["daftar_saham"]):
        listing["kode"] = listing[find_col(listing, "kode")].astype(str).str.upper().str.strip()
        listing["papan pencatatan"] = listing[find_col(listing, "papan pencatatan")].astype(str).str.strip()
        listing["tanggal pencatatan"] = listing[find_col(listing, "tanggal pencatatan")]
    elif has_columns(listing, ALTERNATE_COLUMNS["daftar_saham"]):
        listing["kode"] = listing[find_col(listing, "id instrument")].astype(str).str.upper().str.strip()
        listing["papan pencatatan"] = listing[find_col(listing, "id board")].astype(str).str.strip()
        listing["tanggal pencatatan"] = pd.NaT
        listing = listing.groupby("kode", dropna=False).agg({
            "papan pencatatan": lambda values: ",".join(sorted(set(value for value in values if pd.notna(value)))),
            "tanggal pencatatan": "first",
        }).reset_index()
    else:
        listing = pd.DataFrame({
            "kode": data["kode saham"],
            "papan pencatatan": pd.NA,
            "tanggal pencatatan": pd.NaT,
        })
    data = data.merge(listing[["kode", "papan pencatatan", "tanggal pencatatan"]],
                      left_on="kode saham", right_on="kode", how="left")

    trade = trade_df.copy()
    trade["id instrument"] = trade[find_col(trade, "id instrument")].astype(str).str.upper().str.strip()
    trade["trade_nilai"] = safe_col(trade, "nilai")
    trade["trade_frekuensi"] = safe_col(trade, "frekuensi")
    trade_agg = trade.groupby("id instrument", dropna=False)[["trade_nilai", "trade_frekuensi"]].sum().reset_index()
    data = data.merge(trade_agg.rename(columns={"id instrument": "kode saham"}), on="kode saham", how="left")

    broker = broker_df.copy()
    broker["kode perusahaan"] = broker[find_col(broker, "kode perusahaan")].astype(str).str.upper().str.strip()
    broker["broker_nilai"] = safe_col(broker, "nilai")
    broker["broker_frekuensi"] = safe_col(broker, "frekuensi")
    overlap = len(set(data["kode saham"]) & set(broker["kode perusahaan"]))
    if overlap / max(1, min(len(data), len(broker))) >= 0.2:
        broker_agg = broker.groupby("kode perusahaan", dropna=False)[["broker_nilai", "broker_frekuensi"]].sum().reset_index()
        data = data.merge(broker_agg.rename(columns={"kode perusahaan": "kode saham"}), on="kode saham", how="left")
        data["broker_score"] = percentile_series(data["broker_nilai"].fillna(0)) * 0.6 + percentile_series(data["broker_frekuensi"].fillna(0)) * 0.4
    else:
        data["broker_score"] = 50.0

    data["momentum_score"] = percentile_series(data["change_pct"])
    data["liquidity_score"] = percentile_series(data["nilai_num"]) * 0.5 + percentile_series(data["volume_num"]) * 0.25 + percentile_series(data["frekuensi_num"]) * 0.25
    data["flow_score"] = percentile_series(data["foreign_net_ratio"]) * 0.6 + percentile_series(data["bid_offer_pressure"]) * 0.4
    data["market_activity_score"] = percentile_series(data["trade_nilai"].fillna(0)) * 0.6 + percentile_series(data["trade_frekuensi"].fillna(0)) * 0.4
    data["vol_per_freq"] = np.where(data["frekuensi_num"] != 0, data["volume_num"] / data["frekuensi_num"], np.nan)
    data["volume_trend_score"] = percentile_series(data["vol_per_freq"]) * 0.5 + percentile_series(data["bid_offer_pressure"]) * 0.5
    data["price_structure_score"] = percentile_series(data["close_position"])
    data["final_score"] = sum([
        data["momentum_score"] * weights["momentum"],
        data["liquidity_score"] * weights["liquidity"],
        data["flow_score"] * weights["flow"],
        data["market_activity_score"] * weights["market_activity"],
        data["volume_trend_score"] * weights["volume_trend"],
        data["price_structure_score"] * weights["price_structure"],
        data["broker_score"] * weights["broker"],
    ])
    data["kategori"] = pd.cut(
        data["final_score"],
        bins=[-np.inf, 40, 60, 75, np.inf],
        labels=["Rendah", "Menarik", "Tinggi", "Sangat Tinggi"],
    )
    columns = [
        "kode saham", "nama perusahaan", "papan pencatatan", "tanggal pencatatan",
        "trading_date",
        "penutupan_num", "open_num", "high_num", "low_num", "prev_num",
        "change_pct", "volume_num", "nilai_num", "frekuensi_num",
        "foreign_net", "foreign_net_ratio", "bid_offer_pressure", "atr_pct",
        "close_position", "final_score", "kategori", "momentum_score",
        "liquidity_score", "flow_score", "market_activity_score",
        "volume_trend_score", "price_structure_score", "broker_score",
    ]
    return data[columns].sort_values("final_score", ascending=False).reset_index(drop=True)


def compute_multiday_signals(scored_days, day_labels):
    if len(scored_days) == 1:
        result = scored_days[0].copy()
        result["trend_slope"] = 0.0
        result["score_consistency"] = 0.0
        result["trend_score_norm"] = result["final_score"]
        result["signal_strength"] = result["final_score"]
        result["signal_label"] = result["kategori"].astype(str)
        result["days_data"] = 1
        result["score_day_labels"] = day_labels[0] if day_labels else "D1"
        return result

    panel_rows = []
    for index, (day, label) in enumerate(zip(scored_days, day_labels)):
        part = day[["kode saham", "final_score", "foreign_net_ratio", "nilai_num"]].copy()
        part["day_idx"] = index
        panel_rows.append(part)
    panel = pd.concat(panel_rows, ignore_index=True)
    latest = scored_days[-1].copy()

    def slope_metrics(group):
        ordered = group.sort_values("day_idx")
        scores = ordered["final_score"].values
        if len(scores) < 2:
            return pd.Series({"trend_slope": 0.0, "score_consistency": 0.0, "foreign_acc_mean": 0.0, "vol_growth": 0.0})
        slope = float(np.polyfit(np.arange(len(scores), dtype=float), scores, 1)[0])
        consistency = float(np.mean(scores >= 60) * 100)
        foreign = ordered["foreign_net_ratio"].values
        values = ordered["nilai_num"].values
        growth = float((values[-1] - values[0]) / (values[0] + 1e-9) * 100)
        return pd.Series({
            "trend_slope": slope,
            "score_consistency": consistency,
            "foreign_acc_mean": float(np.mean(foreign)),
            "vol_growth": growth,
        })

    trends = panel.groupby("kode saham").apply(slope_metrics).reset_index()
    result = latest.merge(trends, on="kode saham", how="left")
    result["trend_score_norm"] = np.clip(50 + result["trend_slope"] * 10, 0, 100)
    result["foreign_acc_norm"] = np.clip(50 + result["foreign_acc_mean"] * 500, 0, 100)
    result["vol_growth_norm"] = np.clip(50 + result["vol_growth"] / 4, 0, 100)
    result["signal_strength"] = (
        result["final_score"] * 0.35
        + result["trend_score_norm"] * 0.30
        + result["score_consistency"] * 0.20
        + result["foreign_acc_norm"] * 0.10
        + result["vol_growth_norm"] * 0.05
    )

    def label(row):
        final_score = row["final_score"]
        trend_score = row.get("trend_score_norm", 50)
        consistency = row.get("score_consistency", 0)
        if final_score >= 70 and trend_score >= 60 and consistency >= 60:
            return "🔥 Breakout Candidate"
        if final_score >= 60 and trend_score >= 60:
            return "📈 Trending Up"
        if final_score >= 60 and trend_score < 45:
            return "⚠️ Fading Momentum"
        if final_score < 50 and trend_score >= 60:
            return "👀 Emerging"
        if final_score >= 75:
            return "✅ Sangat Tinggi"
        if final_score >= 60:
            return "✅ Tinggi"
        return "⬜ Watch Only"

    result["signal_label"] = result.apply(label, axis=1)
    result["days_data"] = len(scored_days)
    result["score_day_labels"] = " -> ".join(day_labels)
    return result.sort_values("signal_strength", ascending=False).reset_index(drop=True)


def compute_market_regime(index_df):
    data = index_df.copy()
    data["kode indeks"] = data[find_col(data, "kode indeks")].astype(str).str.upper().str.strip()
    data["sebelumnya_num"] = safe_col(data, "sebelumnya")
    data["selisih_num"] = safe_col(data, "selisih")
    data["nilai_num"] = safe_col(data, "nilai")
    data["frekuensi_num"] = safe_col(data, "frekuensi")
    preferred = data[data["kode indeks"].str.contains("IHSG|COMPOSITE|JCI", regex=True, na=False)]
    row = preferred.iloc[0] if not preferred.empty else data.sort_values("nilai_num", ascending=False).iloc[0]
    previous = safe_num(row["sebelumnya_num"], 0)
    difference = safe_num(row["selisih_num"], 0)
    change = difference / previous * 100 if previous else 0
    momentum = float(np.clip(50 + change * 8, 0, 100))
    activity = percentile_series(data["nilai_num"]).loc[row.name] * 0.6 + percentile_series(data["frekuensi_num"]).loc[row.name] * 0.4
    regime_score = float(np.clip(momentum * 0.7 + activity * 0.3, 0, 100))
    return {
        "index_code": str(row["kode indeks"]),
        "index_change_pct": float(change),
        "regime_score": regime_score,
        "regime_label": "Risk-On" if regime_score >= 60 else "Netral" if regime_score >= 45 else "Risk-Off",
    }


def _ai_row_payload(row):
    """Turn a pandas row into a bounded, JSON-safe AI input."""
    if hasattr(row, "to_dict"):
        row = row.to_dict()
    payload = {}
    for key, value in (row or {}).items():
        if isinstance(value, (np.generic,)):
            value = value.item()
        if isinstance(value, (pd.Timestamp, datetime)):
            value = value.isoformat()
        if isinstance(value, float) and (np.isnan(value) or np.isinf(value)):
            value = None
        try:
            if pd.isna(value):
                value = None
        except (TypeError, ValueError):
            pass
        payload[str(key)] = value
    return payload


def _ai_state_key(ticker):
    return f"structured_ai_analysis_{str(ticker).upper().strip()}"


def render_structured_analysis(analysis, validation_warnings=None, chart_history=None, quant_context=None):
    """Render the stable AI response as UI components instead of raw Markdown."""
    if not isinstance(analysis, dict):
        return
    quant_context = quant_context if isinstance(quant_context, dict) else {}
    quality = analysis.get("data_quality", {})
    structure = analysis.get("market_structure", {})
    levels = analysis.get("levels", {})
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Verdict", str(analysis.get("verdict", "N/A")))
    c2.metric("Confidence", f"{safe_num(analysis.get('confidence'), 0):.0f}/100")
    c3.metric("Data quality", str(quality.get("status", "N/A")))
    c4.metric("Ticker", str(analysis.get("ticker", "N/A")))
    st.info(str(analysis.get("summary", "Belum ada ringkasan.")))
    if validation_warnings:
        st.warning("; ".join(validation_warnings))

    left, right = st.columns(2)
    with left:
        st.markdown("**Market structure**")
        st.write({
            "Trend": structure.get("trend", "N/A"),
            "Momentum": structure.get("momentum", "N/A"),
            "Flow": structure.get("flow", "N/A"),
            "Technical": structure.get("technical", "N/A"),
        })
    with right:
        st.markdown("**Level terverifikasi**")
        level_rows = []
        for key in ["entry_low", "entry_high", "stop", "tp1", "tp2"]:
            if key in levels:
                level_rows.append({"Level": key, "Harga": levels.get(key)})
        if level_rows:
            st.dataframe(pd.DataFrame(level_rows), hide_index=True, width="stretch")
        st.caption(f"Support: {levels.get('support', [])} | Resistance: {levels.get('resistance', [])}")

    st.markdown("**AI Quant Chart**")
    quant_history = pd.DataFrame(quant_context.get("history", []) or [])
    if quant_history.empty and isinstance(chart_history, pd.DataFrame):
        quant_history = chart_history.copy()
    if quant_history.empty:
        st.info("AI Quant Chart belum tersedia karena histori OHLC belum tersedia atau belum cukup.")
    else:
        st.caption("Chart ini menggabungkan OHLC, volume, EMA, support/resistance, Fibonacci, level eksekusi, dan drawing AI yang sudah lolos validasi.")
        try:
            st.plotly_chart(
                create_ai_chart(quant_history, quant_context, analysis),
                use_container_width=True,
                key=f"ai_quant_chart_{str(analysis.get('ticker', 'stock')).upper()}",
            )
        except (RuntimeError, ValueError) as exc:
            st.warning(f"AI Quant Chart gagal dibuat: {exc}")

        indicators = quant_context.get("indicators", {}) or {}
        quant_metrics = st.columns(6)
        quant_metrics[0].metric("Candle", str(quant_context.get("observations", len(quant_history))))
        quant_metrics[1].metric("Close", price_text(indicators.get("close")))
        quant_metrics[2].metric("RSI14", price_text(indicators.get("rsi14")))
        quant_metrics[3].metric("ATR14", price_text(indicators.get("atr14")))
        quant_metrics[4].metric("MACD", price_text(indicators.get("macd")))
        quant_metrics[5].metric("Status", str(quant_context.get("status", "N/A")))

    with st.expander("Detail Quant Engine dan Drawing AI", expanded=True):
        quant_detail_left, quant_detail_right = st.columns(2)
        with quant_detail_left:
            st.markdown("**Indikator terhitung**")
            indicator_rows = [{"Indikator": key, "Nilai": value} for key, value in (quant_context.get("indicators", {}) or {}).items()]
            if indicator_rows:
                st.dataframe(pd.DataFrame(indicator_rows), hide_index=True, width="stretch")
            else:
                st.write("Belum tersedia.")
        with quant_detail_right:
            st.markdown("**Level quant dan eksekusi**")
            quant_level_rows = []
            for group in ("support", "resistance"):
                quant_level_rows.extend({"Kelompok": group, "Harga": value} for value in (quant_context.get("levels", {}).get(group, []) or []))
            quant_level_rows.extend(
                {"Kelompok": key, "Harga": value}
                for key, value in (quant_context.get("execution", {}) or {}).items()
                if key != "atr_proxy"
            )
            if quant_level_rows:
                st.dataframe(pd.DataFrame(quant_level_rows), hide_index=True, width="stretch")
            else:
                st.write("Belum tersedia.")
        drawings = analysis.get("drawings", []) or []
        st.markdown("**Drawing AI tervalidasi**")
        if drawings:
            st.dataframe(
                pd.DataFrame([
                    {
                        "Nama": item.get("name", ""),
                        "Tipe": item.get("type", ""),
                        "Tujuan": item.get("purpose", ""),
                        "Evidence": item.get("evidence", ""),
                    }
                    for item in drawings if isinstance(item, dict)
                ]),
                hide_index=True,
                width="stretch",
            )
        else:
            st.write("Tidak ada drawing AI.")

    st.markdown("**Skenario**")
    scenarios = analysis.get("scenarios", [])
    scenario_rows = [
        {
            "Skenario": item.get("name"),
            "Probabilitas": f"{safe_num(item.get('probability'), 0):.0f}%",
            "Kondisi": item.get("condition", ""),
            "Trigger": item.get("trigger", ""),
            "Invalidasi": item.get("invalidation", ""),
            "Aksi": item.get("action", ""),
        }
        for item in scenarios if isinstance(item, dict)
    ]
    if scenario_rows:
        st.dataframe(pd.DataFrame(scenario_rows), hide_index=True, width="stretch")

    with st.expander("Evidence, risiko, dan trigger", expanded=True):
        for label, values in [
            ("Evidence", structure.get("evidence", [])),
            ("Risiko", analysis.get("risk_flags", [])),
            ("Trigger pemantauan", analysis.get("monitoring_triggers", [])),
            ("Data yang belum tersedia", quality.get("missing", [])),
            ("Konflik data", quality.get("conflicts", [])),
        ]:
            st.markdown(f"**{label}**")
            if values:
                for value in values:
                    st.write(f"- {value}")
            else:
                st.write("Tidak ada.")
    st.markdown("**Kesimpulan**")
    st.write(analysis.get("conclusion", "Belum ada kesimpulan."))
    st.caption(analysis.get("disclaimer", "AI adalah alat bantu analisis, bukan rekomendasi investasi."))
    st.download_button(
        "Download detail analisis Excel",
        build_analysis_excel_bytes(analysis, quant_context, chart_history),
        f"ai_analysis_{str(analysis.get('ticker', 'stock')).upper()}.xlsx",
        EXCEL_MIME,
        key=f"analysis_excel_{str(analysis.get('ticker', 'stock')).upper()}",
        help="Berisi ringkasan, quant indicators, levels, skenario, evidence, drawing, dan data AI Quant Chart.",
    )


def render_advanced_stock_page(openrouter_key, llm_model):
    st.title("BEI Screener — Advanced")
    st.caption("Analisis dapat memakai data tersimpan di Supabase atau upload file BEI baru.")
    database_configured = supabase_is_configured()
    with st.sidebar:
        st.header("Sumber Data Advanced")
        source_options = ["Database Supabase", "Upload file baru"]
        source_index = 0 if database_configured else 1
        data_source = st.radio("Gunakan data dari", source_options, index=source_index, key="advanced_data_source")
        selected_db_dates = []
        day_data = []
        index_file = None
        database_dates = []
        if data_source == "Database Supabase":
            database_dates, dates_error = load_import_dates()
            if dates_error:
                st.error(f"Supabase tidak bisa dibaca: {dates_error}")
            elif database_dates:
                selected_db_dates = st.multiselect(
                    "Pilih tanggal dari database (maksimal 5)",
                    database_dates,
                    default=database_dates[:1],
                    key="advanced_database_dates",
                )
                if len(selected_db_dates) > 5:
                    st.warning("Maksimal 5 tanggal untuk analisis multi-hari.")
            else:
                st.info("Belum ada batch data di Supabase. Pilih Upload file baru.")
        else:
            st.header("Upload Data BEI")
            number_of_days = st.number_input("Jumlah hari", 1, 5, 1, 1)
            for index in range(int(number_of_days)):
                with st.expander(f"Hari {index + 1}", expanded=index == int(number_of_days) - 1):
                    trading_date = st.date_input(
                        "1. Pilih tanggal perdagangan",
                        value=date.today() - timedelta(days=index),
                        key=f"trading_date_{index}",
                    )
                    st.caption("2. Upload file resmi untuk tanggal tersebut")
                    stock_file = st.file_uploader("Ringkasan Saham", type=["xlsx", "xls", "csv"], key=f"s_{index}")
                    broker_file = st.file_uploader("Ringkasan Broker", type=["xlsx", "xls", "csv"], key=f"b_{index}")
                    trade_file = st.file_uploader("Ringkasan Perdagangan", type=["xlsx", "xls", "csv"], key=f"p_{index}")
                    listing_file = st.file_uploader("Daftar Saham", type=["xlsx", "xls", "csv"], key=f"d_{index}")
                    day_data.append({
                        "trading_date": trading_date.isoformat(),
                        "saham": stock_file,
                        "broker": broker_file,
                        "perdagangan": trade_file,
                        "daftar": listing_file,
                    })
            index_file = st.file_uploader("Ringkasan Indeks (Opsional)", type=["xlsx", "xls", "csv"])
        st.markdown("---")
        st.subheader("Bobot Scoring")
        weight_values = {
            "momentum": st.slider("Momentum", 0, 100, 25),
            "liquidity": st.slider("Likuiditas", 0, 100, 20),
            "flow": st.slider("Flow", 0, 100, 20),
            "market_activity": st.slider("Market Activity", 0, 100, 15),
            "volume_trend": st.slider("Volume Trend", 0, 100, 10),
            "price_structure": st.slider("Price Structure", 0, 100, 5),
            "broker": st.slider("Broker", 0, 100, 5),
        }
        total_weight = max(sum(weight_values.values()), 1)
        weights = {key: value / total_weight for key, value in weight_values.items()}

    loaded_batches = []
    if data_source == "Database Supabase":
        if not database_dates or not selected_db_dates:
            st.info("Pilih minimal satu tanggal dari database Supabase.")
            return
        if len(selected_db_dates) > 5:
            return
        with st.spinner("Membaca data BEI dari Supabase..."):
            for trading_date in selected_db_dates:
                frames, load_error = load_import_bundle(trading_date)
                if load_error:
                    st.error(f"[{trading_date}] Gagal membaca database: {load_error}")
                    continue
                if frames:
                    loaded_batches.append({"trading_date": trading_date, "frames": frames, "files": {}})
    else:
        selected_dates = [day["trading_date"] for day in day_data]
        if len(selected_dates) != len(set(selected_dates)):
            st.error("Tanggal upload tidak boleh duplikat. Gunakan satu batch untuk satu tanggal perdagangan.")
            return
        complete_days = [
            day for day in day_data
            if all([day["saham"], day["broker"], day["perdagangan"], day["daftar"]])
        ]
        if not complete_days:
            st.info("Upload minimal 1 set lengkap yang terdiri dari 4 file.")
            return
        with st.spinner("Membaca file BEI..."):
            for day in complete_days:
                try:
                    loaded_batches.append({
                        "trading_date": day["trading_date"],
                        "frames": {
                            "ringkasan_saham": load_table(day["saham"]),
                            "ringkasan_broker": load_table(day["broker"]),
                            "ringkasan_perdagangan": load_table(day["perdagangan"]),
                            "daftar_saham": load_table(day["daftar"]),
                        },
                        "files": {
                            "ringkasan_saham": uploaded_file_payload(day["saham"]),
                            "ringkasan_broker": uploaded_file_payload(day["broker"]),
                            "ringkasan_perdagangan": uploaded_file_payload(day["perdagangan"]),
                            "daftar_saham": uploaded_file_payload(day["daftar"]),
                        },
                    })
                except Exception as exc:
                    st.error(f"[{day['trading_date']}] Gagal: {exc}")
            if index_file and loaded_batches:
                try:
                    index_df = load_table(index_file)
                    valid, missing = validate_columns(index_df, INDEX_REQUIRED_COLUMNS)
                    if valid:
                        loaded_batches[-1]["frames"]["ringkasan_indeks"] = index_df
                        loaded_batches[-1]["files"]["ringkasan_indeks"] = uploaded_file_payload(index_file)
                    else:
                        st.warning(f"Indeks kolom kurang: {', '.join(missing)}")
                except Exception as exc:
                    st.warning(f"Indeks error: {exc}")

    if not loaded_batches:
        st.error("Tidak ada data yang bisa dianalisis.")
        return

    scored_days, labels, market_regime = [], [], None
    with st.spinner("Memvalidasi dan menghitung scoring BEI..."):
        for batch in loaded_batches:
            trading_date = batch["trading_date"]
            frames = batch["frames"]
            required_frames = [
                ("ringkasan_saham", frames.get("ringkasan_saham")),
                ("ringkasan_broker", frames.get("ringkasan_broker")),
                ("ringkasan_perdagangan", frames.get("ringkasan_perdagangan")),
                ("daftar_saham", frames.get("daftar_saham")),
            ]
            valid = all(frame is not None for _, frame in required_frames)
            for name, frame in required_frames:
                if frame is None:
                    st.error(f"[{trading_date}] Dataset {name} tidak ditemukan di database.")
                    continue
                ok, missing = validate_columns(frame, REQUIRED_COLUMNS[name])
                if name in ALTERNATE_COLUMNS and not ok:
                    ok = has_columns(frame, ALTERNATE_COLUMNS[name])
                if not ok:
                    valid = False
                    st.error(f"[{trading_date}] {name}: kurang {', '.join(missing)}")
            if not valid:
                continue
            try:
                scored_days.append(compute_scores(
                    frames["ringkasan_saham"],
                    frames["ringkasan_broker"],
                    frames["ringkasan_perdagangan"],
                    frames["daftar_saham"],
                    weights,
                ))
                labels.append(trading_date)
                index_df = frames.get("ringkasan_indeks")
                if index_df is not None:
                    index_valid, _ = validate_columns(index_df, INDEX_REQUIRED_COLUMNS)
                    if index_valid:
                        market_regime = compute_market_regime(index_df)
            except Exception as exc:
                st.error(f"[{trading_date}] Gagal menghitung scoring: {exc}")

    if not scored_days:
        st.error("Tidak ada data valid.")
        return

    if data_source == "Upload file baru":
        st.subheader("Penyimpanan Supabase")
        if supabase_is_configured():
            st.caption(f"Credential Supabase terdeteksi. Data belum disimpan sampai tombol ditekan. Migrasi: {IMPORT_MIGRATION_FILE}")
            save_key = "save_bei_import_" + "_".join(batch["trading_date"] for batch in loaded_batches)
            if st.button("Simpan batch upload ke Supabase", type="primary", width="stretch", key=save_key):
                for batch in loaded_batches:
                    result = save_bei_import(batch["trading_date"], batch["frames"], batch["files"])
                    if result.get("ok"):
                        st.success(f"{result['message']} Baris: {result.get('row_counts', {})}. File: {result.get('raw_status')}.")
                    else:
                        st.error(result.get("message", "Gagal menyimpan ke Supabase."))
                        if result.get("detail"):
                            st.caption(result["detail"])
        else:
            st.warning("Supabase belum terhubung. Isi SUPABASE_URL dan SUPABASE_SERVICE_ROLE_KEY/SUPABASE_SECRET_KEY, lalu jalankan migrasi SQL.")
    else:
        st.success(f"Sumber analisis: Supabase Database — {', '.join(labels)}")

    combined = compute_multiday_signals(scored_days, labels)
    if market_regime:
        combined["final_score"] = combined["final_score"] * 0.90 + market_regime["regime_score"] * 0.10
        combined["signal_strength"] = combined["signal_strength"] * 0.90 + market_regime["regime_score"] * 0.10

    if market_regime:
        r1, r2, r3, r4 = st.columns(4)
        r1.metric("Index", market_regime["index_code"])
        r2.metric("IHSG", f"{market_regime['index_change_pct']:+.2f}%")
        r3.metric("Regime", market_regime["regime_label"])
        r4.metric("Regime Score", f"{market_regime['regime_score']:.1f}")

    min_signal = st.slider("Min Signal Strength", 0, 100, 60)
    boards = sorted([str(value) for value in combined["papan pencatatan"].dropna().unique() if str(value).strip()])
    board_filter = st.selectbox("Papan", ["Semua"] + boards)
    top_n = st.number_input("Top N", 5, 300, 25, 5)
    signal_filter = st.selectbox("Sinyal", ["Semua", "Breakout Candidate", "Trending Up", "Fading Momentum", "Emerging"])
    view = combined[combined["signal_strength"] >= min_signal].copy()
    if board_filter != "Semua":
        view = view[view["papan pencatatan"] == board_filter]
    if signal_filter != "Semua":
        view = view[view["signal_label"].str.contains(signal_filter, na=False)]
    view = view.head(int(top_n))

    st.subheader("Kandidat Saham")
    columns = [
        "kode saham", "nama perusahaan", "papan pencatatan", "penutupan_num",
        "change_pct", "foreign_net", "final_score", "signal_strength", "signal_label",
    ]
    if len(scored_days) > 1:
        columns += ["trend_slope", "score_consistency"]
    display = view[[column for column in columns if column in view.columns]].rename(columns={
        "penutupan_num": "close",
        "foreign_net": "net_foreign",
        "signal_strength": "signal",
        "signal_label": "sinyal",
        "trend_slope": "tren/hari",
        "score_consistency": "konsistensi%",
    })
    render_df_with_style_fallback(display, [column for column in ["signal", "final_score"] if column in display.columns])
    st.download_button(
        "Export Excel",
        dataframe_to_excel_bytes({"Kandidat Advanced": view}),
        "bei_advanced.xlsx",
        EXCEL_MIME,
        help="Workbook Excel berisi semua row dan kolom hasil scoring; header sudah freeze dan filter aktif.",
    )
    if view.empty:
        st.warning("Tidak ada kandidat.")
        return

    selected_code = st.selectbox("Pilih saham untuk analisis", view["kode saham"].tolist(), key="advanced_selected")
    selected = combined[combined["kode saham"] == selected_code].iloc[0]
    company = str(selected.get("nama perusahaan", ""))
    st.subheader(f"{selected_code}  {company}  {selected.get('signal_label', '')}")
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Signal", f"{selected.get('signal_strength', selected['final_score']):.1f}")
    m2.metric("Score", f"{selected['final_score']:.1f}")
    m3.metric("Change %", f"{selected['change_pct']:+.2f}%")
    m4.metric("Net Foreign", f"{selected['foreign_net']:,.0f}")
    m5.metric("ATR-1d", f"{selected.get('atr_pct', 0):.2f}%")

    tab_score, tab_chart, tab_ai = st.tabs(["Score", "Chart", "AI Analisis"])
    with tab_score:
        breakdown = pd.DataFrame({
            "Faktor": ["Momentum", "Likuiditas", "Flow", "Market Activity", "Vol Trend", "Price Structure", "Broker"],
            "Skor": [safe_num(selected.get(column), np.nan) for column in [
                "momentum_score", "liquidity_score", "flow_score", "market_activity_score",
                "volume_trend_score", "price_structure_score", "broker_score",
            ]],
        })
        breakdown["Status"] = breakdown["Skor"].apply(factor_label)
        render_df_with_style_fallback(breakdown, ["Skor"])
        if len(scored_days) > 1:
            x1, x2, x3 = st.columns(3)
            x1.metric("Tren/Hari", f"{selected.get('trend_slope', 0):+.1f} pts")
            x2.metric("Konsistensi", f"{selected.get('score_consistency', 0):.0f}%")
            x3.metric("Hari Data", f"{selected.get('days_data', 1)}")
        close = safe_num(selected.get("penutupan_num"), np.nan)
        atr_pct = max(safe_num(selected.get("atr_pct"), 2.0), 0.5)
        if not pd.isna(close):
            atr_abs = close * atr_pct / 100
            st.markdown("**Level Eksekusi**")
            levels = st.columns(5)
            levels[0].metric("Entry Low", f"{close * 0.995:.2f}")
            levels[1].metric("Entry High", f"{close + atr_abs * 0.3:.2f}")
            levels[2].metric("Stop", f"{close - atr_abs * 1.5:.2f}")
            levels[3].metric("TP1", f"{close + atr_abs * 2.0:.2f}")
            levels[4].metric("TP2", f"{close + atr_abs * 3.5:.2f}")
            st.caption(f"R/R: {(atr_abs * 2.0) / (atr_abs * 1.5):.2f}x | ATR={atr_abs:.2f} ({atr_pct:.2f}%)")

    with tab_chart:
        tv_symbol = f"IDX:{selected_code}"
        tv_slug = f"IDX-{selected_code}"
        history = build_price_history(scored_days, selected_code, labels)
        quant_context = build_quant_context(history, _ai_row_payload(selected))
        ai_result = st.session_state.get(_ai_state_key(selected_code))
        chart_tab, tradingview_tab, technical_tab, financial_tab, news_tab = st.tabs(
            ["AI Chart", "TradingView", "Technicals", "Financials", "News"]
        )
        with chart_tab:
            st.caption("Plotly memakai OHLC BEI yang diunggah. Garis AI hanya muncul setelah respons JSON lolos validasi.")
            try:
                chart_history = pd.DataFrame(quant_context.get("history", []))
                if chart_history.empty:
                    chart_history = history
                st.plotly_chart(create_ai_chart(chart_history, quant_context, ai_result), use_container_width=True)
            except (RuntimeError, ValueError) as exc:
                st.warning(str(exc))
            if quant_context.get("observations", 0) < 14:
                st.info("Histori belum cukup untuk indikator periode panjang. Level awal memakai range yang tersedia dan ditandai sebagai proxy.")
            st.json({
                "observations": quant_context.get("observations"),
                "indicator_status": quant_context.get("status"),
                "support": quant_context.get("levels", {}).get("support", []),
                "resistance": quant_context.get("levels", {}).get("resistance", []),
                "fibonacci": quant_context.get("levels", {}).get("fibonacci", []),
            })
        with tradingview_tab:
            render_tradingview_advanced_chart(tv_symbol, interval="D", height=580)
        with technical_tab:
            components.html(
                f"""<div class="tradingview-widget-container" style="width:100%;height:700px;">
<div class="tradingview-widget-container__widget" style="width:100%;height:100%;"></div>
<script src="https://s3.tradingview.com/external-embedding/embed-widget-technical-analysis.js" async>
{json.dumps({"interval": "1D", "width": "100%", "height": 700, "symbol": tv_symbol, "showIntervalTabs": True, "locale": "id", "colorTheme": "dark"})}
</script></div>""",
                height=720,
            )
        with financial_tab:
            components.iframe(f"https://id.tradingview.com/symbols/{tv_slug}/financials-overview/", height=700, scrolling=True)
        with news_tab:
            components.html(
                f"""<div class="tradingview-widget-container" style="width:100%;height:700px;">
<div class="tradingview-widget-container__widget" style="width:100%;height:100%;"></div>
<script src="https://s3.tradingview.com/external-embedding/embed-widget-timeline.js" async>
{json.dumps({"feedMode": "symbol", "symbol": tv_symbol, "displayMode": "regular", "width": "100%", "height": 700, "colorTheme": "dark", "locale": "id", "isTransparent": False})}
</script></div>""",
                height=720,
            )

    with tab_ai:
        analysis_rendered = False
        if not openrouter_key:
            st.warning("Masukkan OpenRouter API Key di sidebar.")
        elif st.button("Generate Analisis AI", width="stretch", key=f"advanced_ai_{selected_code}"):
            with st.spinner("TradingView dan OpenRouter menganalisis saham..."):
                tv_raw = fetch_tv_data(selected_code)
                tv_parsed = parse_tv_data(tv_raw)
                if tv_parsed:
                    st.success("Data teknikal TradingView berhasil dimuat.")
                    with st.expander("Data Teknikal TradingView", expanded=True):
                        rec_all = tv_parsed.get("rec_all", "N/A")
                        a1, a2, a3 = st.columns(3)
                        a1.metric("Rekomendasi", f"{tv_rec_emoji(rec_all)} {rec_all}")
                        a2.metric("MA Signal", f"{tv_rec_emoji(tv_parsed.get('rec_ma', ''))} {tv_parsed.get('rec_ma', 'N/A')}")
                        a3.metric("Oscillator", f"{tv_rec_emoji(tv_parsed.get('rec_other', ''))} {tv_parsed.get('rec_other', 'N/A')}")
                else:
                    st.warning("Data TradingView tidak tersedia. AI memakai data BEI saja.")
                context = None
                if len(scored_days) > 1:
                    context = {
                        "trend_slope": safe_num(selected.get("trend_slope"), 0),
                        "score_consistency": safe_num(selected.get("score_consistency"), 0),
                        "days_data": int(safe_num(selected.get("days_data"), 1)),
                    }
                runtime_context = build_runtime_context(
                    selected_code,
                    company,
                    _ai_row_payload(selected),
                    tv_parsed,
                    quant_context,
                    market_regime,
                    context,
                )
                outcome = call_openrouter_structured(runtime_context, openrouter_key, llm_model)
                if not outcome.get("ok"):
                    st.error(outcome.get("error", "Analisis AI gagal."))
                else:
                    try:
                        validated, validation_warnings = validate_analysis(
                            outcome["data"], selected_code, history, quant_context
                        )
                        st.session_state[_ai_state_key(selected_code)] = validated
                        save_status = save_analysis(selected_code, runtime_context, validated)
                        st.success("Analisis AI selesai dan drawing sudah divalidasi.")
                        st.caption(f"Storage: lokal={save_status['local']} | Supabase={save_status['supabase']}")
                        render_structured_analysis(
                            validated,
                            validation_warnings,
                            chart_history=history,
                            quant_context=quant_context,
                        )
                        analysis_rendered = True
                    except ValueError as exc:
                        st.error(f"Validasi respons AI gagal: {exc}")
        stored_analysis = st.session_state.get(_ai_state_key(selected_code))
        if stored_analysis and not analysis_rendered:
            st.caption("Memuat hasil analisis terakhir pada sesi ini.")
            render_structured_analysis(
                stored_analysis,
                stored_analysis.get("validation_warnings", []),
                chart_history=history,
                quant_context=quant_context,
            )


# ── NAVIGATION ────────────────────────────────────────────────────────────────
if "page" not in st.session_state:
    st.session_state["page"] = "Home"
if st.session_state["page"] not in APP_PAGES:
    st.session_state["page"] = "Home"

with st.sidebar:
    st.header("Navigasi")
    nav_choice = st.selectbox(
        "Mode",
        APP_PAGES,
        index=APP_PAGES.index(st.session_state["page"]),
        key=f"nav_mode_{st.session_state['page'].replace(' ', '_')}",
    )
    if nav_choice != st.session_state["page"]:
        st.session_state["page"] = nav_choice
        st.rerun()
    st.markdown("---")

if st.session_state["page"] == "Home":
    render_home()
    st.stop()
if st.session_state["page"] == "Watchlist & Alerts":
    render_watchlist_page()
    st.stop()

st.title("BEI Stock Screener")
st.caption("Auto TradingView untuk screening cepat atau Upload BEI Advanced untuk flow detail.")
with st.sidebar:
    openrouter_key, llm_model = render_openrouter_sidebar()
    st.markdown("---")
    stock_mode = st.radio(
        "Sumber Data Saham",
        ["Auto TradingView", "Upload BEI Advanced"],
        index=0,
        key="stock_data_mode",
    )

if stock_mode == "Auto TradingView":
    render_stock_auto_page(openrouter_key, llm_model)
    st.stop()

render_advanced_stock_page(openrouter_key, llm_model)
